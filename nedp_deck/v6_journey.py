"""Slide 3 of 6M_V6: rebuild 'Development program' as a journey timeline (text and photos unchanged)."""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
E = lambda v: Emu(round(v * 9525))
BLK, GREY, LG = RGBColor(0, 0, 0), RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0xBF, 0xBF, 0xBF)
P = Presentation(sys.argv[1]); s = P.slides[2]
sh = {x.name: x for x in s.shapes}
def nostyle(x):
    st = x._element.find(qn('p:style'))
    if st is not None: x._element.remove(st)
for x in list(s.shapes):
    if x.shape_type == 9: x._element.getparent().remove(x._element)   # old arrows
def place(x, l, t, w, h): x.left, x.top, x.width, x.height = E(l), E(t), E(w), E(h)
def fsize(x, pt, color=None, bold=None):
    for p in x.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)
            if color is not None: r.font.color.rgb = color
            if bold is not None: r.font.bold = bold
    x.text_frame.word_wrap = True
    from pptx.enum.text import PP_ALIGN
    for p in x.text_frame.paragraphs:
        if p.alignment not in (None, PP_ALIGN.LEFT): p.alignment = PP_ALIGN.LEFT
def box(l, t, w, h, text, pt, color=BLK, bold=False):
    b = s.shapes.add_textbox(E(l), E(t), E(w), E(h)); tf = b.text_frame
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, a, 0)
    tf.word_wrap = True; from pptx.enum.text import MSO_AUTO_SIZE; tf.auto_size = MSO_AUTO_SIZE.NONE
    from pptx.enum.text import PP_ALIGN
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    r = tf.paragraphs[0].add_run(); r.text = text; r.font.size = Pt(pt); r.font.bold = bold; r.font.color.rgb = color
    return b
COLS = [(36, 330), (386, 330), (736, 160), (906, 160), (1076, 168)]
STAT = ['DONE', 'DONE', 'NOW', 'NEXT', 'NEXT']
TITLES = ['TextBox 10', 'TextBox 19', 'TextBox 29', 'TextBox 32', 'TextBox 40']
LY = 214
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(36), E(LY), E(1244), E(LY)); nostyle(ln)
ln.line.color.rgb = BLK; ln.line.width = Pt(1.5)
for i, ((x, w), st, tn) in enumerate(zip(COLS, STAT, TITLES)):
    d = 22 if st == 'NOW' else 16
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(LY - d / 2), E(d), E(d)); nostyle(c)
    c.line.color.rgb = BLK; c.line.width = Pt(1.5)
    if st == 'NEXT': c.fill.solid(); c.fill.fore_color.rgb = RGBColor(255, 255, 255)
    else: c.fill.solid(); c.fill.fore_color.rgb = BLK
    box(x, LY + 18, w, 16, st, 10, BLK if st != 'NEXT' else GREY, True)
    t = sh[tn]; place(t, x - 4, LY + 34, w + 4, 34); fsize(t, 20 if w > 200 else 18, BLK if st != 'NEXT' else GREY)
for nm, (x, w) in (('TextBox 12', COLS[0]), ('TextBox 22', COLS[1])):
    place(sh[nm], x - 4, LY + 70, w, 22); fsize(sh[nm], 12, GREY)
for nm, (x, w) in (('TextBox 15', COLS[0]), ('TextBox 24', COLS[1])):
    place(sh[nm], x - 4, LY + 94, w, 96); fsize(sh[nm], 13, BLK)
place(sh['TextBox 54'], 36, 160, 400, 28)
pw, ph, py = 161, 250, 412
def cover(p, x, y, w, h):
    iw, ih = p.image.size
    p.crop_left = p.crop_right = p.crop_top = p.crop_bottom = 0
    r_box, r_img = w / h, iw / ih
    if r_img > r_box: c = (1 - r_box / r_img) / 2; p.crop_left = p.crop_right = c
    else: c = (1 - r_img / r_box) / 2; p.crop_top = p.crop_bottom = c
    place(p, x, y, w, h)
cover(sh['Picture 42'], 36, py, pw, ph); cover(sh['Picture 44'], 36 + pw + 8, py, pw, ph)
cover(sh['Picture 50'], 386, py, 330, (ph - 8) / 2); cover(sh['Picture 52'], 386, py + (ph - 8) / 2 + 8, 330, (ph - 8) / 2)
P.save(sys.argv[2]); print('ok')
