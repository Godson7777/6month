"""NEDP Mid Year Review deck generator. Builds the 15-slide deck in three design options."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

FONT = "Neue Haas Grotesk Text Pro"
W, H, M = 13.333, 7.5, 0.45

C = dict(orange="D85A30", orange_dk="4A1B0C", orange_tint="FBF3EE", black="1C1B19",
         gray_dk="3A3936", navy="1B3A52", muted="8A8880", line="D9D6CE", pale="F2F0EA",
         white="FFFFFF", text="2C2C2A", text2="6E6C64", card="E0DDD4",
         proc_bg="EDEBE0", proc_ln="B0AC9C")

THEMES = {
    "A_Minimal": dict(name="Option A — Minimal Line", bg="FFFFFF", card_fill="FFFFFF", card_line=C["card"],
                      card_style="border", divider="circle", title_rule=False, side_band=False),
    "B_Bold": dict(name="Option B — Bold Contrast", bg="FFFFFF", card_fill=C["pale"], card_line=None,
                   card_style="fill", divider="dark", title_rule=False, side_band=True),
    "C_Editorial": dict(name="Option C — Warm Editorial", bg="F7F5F0", card_fill="FFFFFF", card_line=None,
                        card_style="topline", divider="number", title_rule=True, side_band=False),
}


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, theme):
        self.t = theme
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.n = 0

    # ---------- primitives ----------
    def rect(self, s, x, y, w, h, fill=None, line=None, lw=0.75, dash=False, shape=MSO_SHAPE.RECTANGLE):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.08
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
            if dash:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        return sh

    def line(self, s, x1, y1, x2, y2, color=C["line"], lw=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = rgb(color); ln.line.width = Pt(lw)
        return ln

    def text(self, s, x, y, w, h, content, size=9, bold=False, color=C["text"], align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=None):
        """content: str, or list of paragraphs; each paragraph a str or list of (text, opts) runs."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.01)
        tf.vertical_anchor = anchor
        paras = content if isinstance(content, list) else [content]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            if spacing:
                para.space_after = Pt(spacing)
            runs = p if isinstance(p, list) else [(p, {})]
            for txt, o in runs:
                r = para.add_run(); r.text = txt
                f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
                f.color.rgb = rgb(o.get("color", color))
        return tb

    def bullets(self, s, x, y, w, h, items, size=8.2, color=C["text"], dot=C["orange"], spacing=2.5):
        return self.text(s, x, y, w, h, [[("•  ", {"color": dot, "bold": True}), (it, {})] for it in items],
                         size=size, color=color, spacing=spacing)

    def card(self, s, x, y, w, h, accent=C["orange"], bar=True):
        t = self.t
        if t["card_style"] == "border":
            self.rect(s, x, y, w, h, fill=t["card_fill"], line=t["card_line"], lw=0.75)
            if bar:
                self.rect(s, x, y, 0.05, h, fill=accent)
        elif t["card_style"] == "fill":
            self.rect(s, x, y, w, h, fill=t["card_fill"])
            if bar:
                self.rect(s, x, y, 0.05, h, fill=accent)
        else:  # topline
            self.rect(s, x, y, w, h, fill=t["card_fill"])
            self.rect(s, x, y, w, 0.04 if bar else 0.015, fill=accent if bar else C["line"])

    def evidence(self, s, x, y, w, h, label="Evidence"):
        self.rect(s, x, y, w, h, fill=None if self.t["bg"] == "FFFFFF" else "FFFFFF", line=C["muted"], lw=0.75, dash=True)
        self.text(s, x, y, w, h, label, size=8, color=C["muted"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def badge_num(self, s, x, y, n, color, d=0.32):
        c = self.rect(s, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL)
        self.text(s, x, y, d, d, str(n), size=9, bold=True, color=C["white"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def pill(self, s, x, y, txt, color, w=None):
        w = w or 0.12 + 0.065 * len(txt)
        self.rect(s, x, y, w, 0.24, fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        self.text(s, x, y, w, 0.24, txt, size=7.5, bold=True, color=C["white"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ---------- slide chrome ----------
    def new(self, dark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6]); self.n += 1
        bg = C["black"] if dark else self.t["bg"]
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
        fg, mut = (C["white"], "A9A69C") if dark else (C["text"], C["muted"])
        self.text(s, 0, 0.12, W, 0.2, "[ INTERNAL USE ONLY ]", size=7, color=mut, align=PP_ALIGN.CENTER)
        bx = W - M - 0.95
        self.rect(s, bx, 0.14, 0.95, 0.26, fill=C["orange"] if dark else C["black"])
        self.text(s, bx, 0.14, 0.95, 0.26, "TRIATRA", size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, M, H - 0.36, 4, 0.2, "NEDP Mid Year Review 2026", size=7.5, color=mut)
        self.text(s, W - M - 1, H - 0.36, 1, 0.2, f"{self.n:02d}", size=7.5, color=mut, align=PP_ALIGN.RIGHT)
        if self.t["side_band"] and not dark:
            self.rect(s, 0, 0, 0.14, H, fill=C["black"])
            self.rect(s, 0, 0.55, 0.14, 0.6, fill=C["orange"])
        return s

    def title(self, s, title, kicker):
        self.text(s, M, 0.52, W - 2 * M - 1.1, 0.5, title, size=15, bold=True, color=C["black"], anchor=MSO_ANCHOR.BOTTOM)
        if self.t["title_rule"]:
            self.rect(s, M + 0.02, 1.08, 0.5, 0.035, fill=C["orange"])
            self.text(s, M + 0.62, 1.0, 6, 0.2, kicker.upper(), size=8.5, bold=True, color=C["orange"])
        else:
            self.text(s, M, 1.04, 8, 0.2, kicker.upper(), size=8.5, bold=True, color=C["orange"])
            self.line(s, M, 1.32, W - M, 1.32, C["line"], 0.5)

    # ---------- slides ----------
    SECTIONS = ["Company Profile", "Site Operation", "Organization Chart", "Roles & Responsibility", "Job Description Execution"]

    def agenda(self):
        s = self.new()
        self.title(s, "Five Sections Prove I Know the Company, My Products and How I Execute", "Agenda")
        y = 1.65
        for i, sec in enumerate(self.SECTIONS):
            yy = y + i * 0.92
            if self.t["card_style"] == "border":
                self.rect(s, M, yy, 7.9, 0.78, fill=C["pale"])
            elif self.t["card_style"] == "fill":
                self.rect(s, M, yy, 7.9, 0.78, fill=C["pale"]); self.rect(s, M, yy, 0.05, 0.78, fill=C["orange"])
            else:
                self.line(s, M, yy + 0.8, M + 7.9, yy + 0.8, C["line"], 0.75)
            self.text(s, M + 0.25, yy, 0.9, 0.78, f"{i+1:02d}", size=22, bold=True, color=C["orange"], anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, M + 1.25, yy, 6, 0.78, sec, size=13, bold=True, color=C["black"], anchor=MSO_ANCHOR.MIDDLE)
        x = 9.0
        self.card(s, x, 1.65, W - M - x, 4.5)
        self.text(s, x + 0.3, 1.9, 3.3, 0.25, "PRESENTER", size=8, bold=True, color=C["orange"])
        self.text(s, x + 0.3, 2.2, 3.3, 0.6, "Tiberias K. Simu", size=16, bold=True, color=C["black"])
        self.text(s, x + 0.3, 2.75, 3.3, 0.5, "Marketing Strategic 2 Associate", size=10, color=C["text2"])
        self.line(s, x + 0.3, 3.35, W - M - 0.3, 3.35, C["line"])
        self.text(s, x + 0.3, 3.5, 3.3, 0.25, "PRODUCTS I OWN", size=8, bold=True, color=C["orange"])
        self.text(s, x + 0.3, 3.8, 3.3, 0.9, [[("Small Suppeq", {"bold": True})], [("Truck Others", {"bold": True})]], size=12, color=C["black"], spacing=4)
        self.line(s, x + 0.3, 4.75, W - M - 0.3, 4.75, C["line"])
        self.text(s, x + 0.3, 4.9, 3.3, 1.0, "PT Triatra Sinergia Pratama\nMarketing Unit & Strategic Division", size=8.5, color=C["text2"])

    def divider(self, num, title, subs=None, active=None):
        t = self.t
        if t["divider"] == "dark":
            s = self.new(dark=True)
            self.text(s, M, 1.2, 6, 2.6, f"{num:02d}", size=150, bold=True, color=C["orange"])
            self.text(s, M, 4.2, 2, 0.3, f"SECTION {num:02d}", size=9, bold=True, color="A9A69C")
            self.text(s, M, 4.5, 9, 0.8, title, size=32, bold=True, color=C["white"])
            self.rect(s, M + 0.02, 5.35, 0.8, 0.05, fill=C["orange"])
            fg, act = "A9A69C", C["orange"]
        elif t["divider"] == "number":
            s = self.new()
            self.text(s, W - 6.3, 0.9, 6, 4.5, f"{num:02d}", size=260, bold=True, color="EBD9CF", align=PP_ALIGN.RIGHT)
            self.text(s, M, 2.7, 2, 0.3, f"SECTION {num:02d}", size=9, bold=True, color=C["orange"])
            self.text(s, M, 3.0, 8, 0.8, title, size=32, bold=True, color=C["black"])
            self.rect(s, M + 0.02, 3.88, 0.8, 0.05, fill=C["orange"])
            fg, act = C["muted"], C["orange"]
        else:
            s = self.new()
            self.rect(s, W - 4.6, -1.6, 6.2, 6.2, fill=C["orange_tint"], shape=MSO_SHAPE.OVAL)
            self.text(s, M, 2.7, 2, 0.3, f"SECTION {num:02d}", size=9, bold=True, color=C["orange"])
            self.text(s, M, 3.0, 8, 0.8, title, size=32, bold=True, color=C["black"])
            self.rect(s, M + 0.02, 3.88, 0.8, 0.05, fill=C["orange"])
            fg, act = C["muted"], C["orange"]
        if subs:
            base = 5.75 if t["divider"] == "dark" else 4.25
            for i, sub in enumerate(subs):
                on = i == active
                self.text(s, M, base + i * 0.36, 8, 0.3, sub, size=11, bold=on, color=act if on else fg)

    def company(self):
        s = self.new()
        self.title(s, "Triatra Is Astra's Mining Support Arm — From Unit Selling to After Sales", "Company Profile")
        chain = [("Astra International", "Diversified conglomerate"), ("United Tractors", "Heavy equipment, mining & energy"),
                 ("UTPE", "United Tractors Pandu Engineering — manufacturing"), ("PT Triatra Sinergia Pratama", "Marketing & sales of Patria products and after-sales services")]
        y = 1.65
        self.text(s, M, y, 4, 0.25, "OWNERSHIP CHAIN", size=8, bold=True, color=C["muted"])
        for i, (n, d) in enumerate(chain):
            yy = y + 0.35 + i * 1.12
            last = i == 3
            if last:
                self.rect(s, M, yy, 5.2, 0.8, fill=C["orange_tint"], line=C["orange"], lw=1.25)
            else:
                self.card(s, M, yy, 5.2, 0.8, accent=C["gray_dk"], bar=False)
                if self.t["card_style"] == "topline":
                    pass
            self.text(s, M + 0.2, yy + 0.1, 4.8, 0.3, n, size=11, bold=True, color=C["orange_dk"] if last else C["black"])
            self.text(s, M + 0.2, yy + 0.42, 4.8, 0.3, d, size=8.2, color=C["orange_dk"] if last else C["text2"])
            if not last:
                self.line(s, M + 2.6, yy + 0.8, M + 2.6, yy + 1.12, C["muted"], 1)
        x = 6.2; w = W - M - x
        self.text(s, x, y, 4, 0.25, "BUSINESS LINES", size=8, bold=True, color=C["muted"])
        bw = (w - 0.2) / 2
        for i, (n, d) in enumerate([("PATRIA", "Unit selling — vessels, trailers, supporting equipment and special-purpose trucks"),
                                    ("ULTRA", "After sales — parts, service and product support for the installed fleet")]):
            xx = x + i * (bw + 0.2)
            self.card(s, xx, y + 0.35, bw, 1.55, accent=C["orange"] if i == 0 else C["gray_dk"])
            self.text(s, xx + 0.25, y + 0.5, bw - 0.4, 0.4, n, size=16, bold=True, color=C["black"])
            self.text(s, xx + 0.25, y + 0.95, bw - 0.4, 0.8, d, size=8.5, color=C["text2"])
        self.text(s, x, y + 2.15, 4, 0.25, "INDUSTRIES SERVED", size=8, bold=True, color=C["muted"])
        inds = ["Coal & Mineral Mining", "Construction", "Agro & Forestry", "Marine"]
        for i, ind in enumerate(inds):
            xx = x + (i % 2) * (bw + 0.2); yy = y + 2.5 + (i // 2) * 1.25
            self.card(s, xx, yy, bw, 1.05, bar=False)
            self.text(s, xx + 0.25, yy + 0.12, 0.6, 0.4, f"{i+1:02d}", size=14, bold=True, color=C["orange"])
            self.text(s, xx + 0.25, yy + 0.55, bw - 0.4, 0.35, ind, size=11, bold=True, color=C["black"])

    def site(self):
        s = self.new()
        self.title(s, "Our Coverage Follows Our Customers — From Head Office to Every Active Mine Site", "Site Operation")
        self.evidence(s, M, 1.6, 9.3, 5.25, "Coverage map — paste here")
        x = M + 9.5; w = W - M - x
        for i, (n, d) in enumerate([("Head Office", "Jakarta — marketing, pricing, coordination"),
                                    ("Site Operation", "[ add site list ]"), ("Site Representative", "[ add site list ]")]):
            yy = 1.6 + i * 1.78
            self.card(s, x, yy, w, 1.6, accent=[C["orange"], C["gray_dk"], C["navy"]][i])
            self.rect(s, x + 0.22, yy + 0.22, 0.16, 0.16, fill=[C["orange"], C["gray_dk"], C["navy"]][i], shape=MSO_SHAPE.OVAL)
            self.text(s, x + 0.48, yy + 0.15, w - 0.6, 0.3, n, size=10, bold=True, color=C["black"])
            self.text(s, x + 0.22, yy + 0.55, w - 0.4, 0.9, d, size=8.2, color=C["text2"])

    def org(self):
        s = self.new()
        self.title(s, "Each Associate Owns a Distinct Product Group — Mine Is Small Suppeq and Truck Others", "Organization Chart")
        cx = W / 2
        def box(x, y, w, h, name, role=None, hl=False, products=None, dashed=False, head=False):
            if hl:
                self.rect(s, x, y, w, h, fill=C["orange_tint"], line=C["orange"], lw=1.5)
            elif dashed:
                self.rect(s, x, y, w, h, fill=None, line=C["muted"], dash=True)
            elif head:
                self.rect(s, x, y, w, h, fill=C["black"] if self.t["card_style"] == "fill" else (C["pale"] if self.t["card_style"] == "border" else "FFFFFF"),
                          line=None if self.t["card_style"] != "border" else C["line"])
            else:
                self.card(s, x, y, w, h, bar=False)
            dark = head and self.t["card_style"] == "fill"
            nc = C["orange_dk"] if hl else (C["white"] if dark else C["black"])
            if products is not None:
                self.text(s, x + 0.12, y + 0.06, w - 0.2, 0.22, name, size=9, bold=True, color=nc)
                self.text(s, x + 0.12, y + 0.27, w - 0.2, h - 0.3, products, size=7.6, color=C["orange_dk"] if hl else C["text2"])
            else:
                runs = [[(name, {"bold": True, "color": C["muted"] if dashed else nc})]]
                if role:
                    runs.append([(role, {"size": 7.4, "color": "C9C6BC" if dark else C["muted"]})])
                self.text(s, x, y, w, h, runs, size=9, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        bw = 3.4
        box(cx - bw / 2, 1.5, bw, 0.42, "Ceisar Centiaga", "Director in Charge", head=True)
        self.line(s, cx, 1.92, cx, 2.05, C["muted"], 1)
        box(cx - bw / 2, 2.05, bw, 0.42, "Ricardo Tanada", "Marketing Unit & Strategic Division Head", head=True)
        gap = 0.2; cw = (W - 2 * M - 3 * gap) / 4
        self.line(s, cx, 2.47, cx, 2.6, C["muted"], 1)
        self.line(s, M + cw / 2, 2.6, M + 3 * (cw + gap) + cw / 2, 2.6, C["muted"], 1)
        cols = [
            ("Kresnanto A.W.", "Marketing Strategic 1 Head", [("Satrio F.R.", "Big Suppeq, Big Vessel, Bucket"),
             ("M. Fawwaz J.", "Cabin, Forestry, Agro, Forklift, Towerlamp, Pump")]),
            ("Ricardo Tanada (c)", "Marketing Strategic 2 Head", None),
            ("Ugik Praseno", "Marketing Strategic 3 Head", [("Baihaqi", "Mineral Processing, EPC"), ("Yuda Prana", "Maritime, Renewable Energy"),
             ("Celine", "Ultra Product Support"), ("Hanifa", "Automation, EV")]),
            ("Marketing Communication", None, [("Karina Audina E.", ""), ("Adinda Siti K.", ""), ("Rahma", "Graphic Designer")]),
        ]
        top = 2.78; ph = 0.56; pg = 0.08
        for i, (hn, hr, ppl) in enumerate(cols):
            x = M + i * (cw + gap)
            self.line(s, x + cw / 2, 2.6, x + cw / 2, top, C["muted"], 1)
            if i == 1:
                self.rect(s, x - 0.08, top - 0.08, cw + 0.16, 3.52, fill=None, line=C["orange"], lw=1.25)
                box(x, top, cw, 0.5, hn, hr, head=True)
                y = top + 0.5 + pg
                box(x, y, cw, 0.42, "Ignatius Indrawan S.", "Deputy Head"); y += 0.42 + pg
                for n, p, hl in [("Ervin Bahar P.", "Medium Vessel, Man Hauler, Bus Scania", False),
                                 ("Tiberias K.S.", "Small Suppeq, Truck Others", True),
                                 ("Eurico", "Mining / Non-mining Trailer, Overseas Zone", False)]:
                    box(x, y, cw, ph, n, products=p, hl=hl); y += ph + pg
                box(x, y, cw, 0.36, "Nandu Narpati — Admin", dashed=True)
                continue
            box(x, top, cw, 0.5, hn, hr, head=True)
            y = top + 0.5 + pg
            for n, p in ppl:
                h = ph if p else 0.36
                if p and i == 3:
                    h = 0.36; box(x, y, cw, h, f"{n} — {p}")
                elif p:
                    box(x, y, cw, h, n, products=p)
                else:
                    box(x, y, cw, h, n)
                y += h + pg

    def rr_a(self):
        s = self.new()
        self.title(s, "I Report Through the Deputy Head and Own Two Product Groups End to End", "Roles & Responsibility · Position & Products")
        self.text(s, M, 1.6, 4, 0.25, "REPORTING LINE", size=8, bold=True, color=C["muted"])
        lw = 5.0; cx = M + lw / 2
        rows = [(lw * 0.62, "Ricardo Tanada", "Dept Head (concurrent)"), (lw * 0.78, "Ignatius Indrawan Sumawinata", "Deputy Head")]
        y = 1.95
        for w, n, r in rows:
            self.card(s, cx - w / 2, y, w, 0.66, bar=False)
            self.text(s, cx - w / 2, y, w, 0.66, [[(n, {"bold": True, "color": C["black"]})], [(r, {"size": 7.6, "color": C["muted"]})]],
                      size=9.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self.line(s, cx, y + 0.66, cx, y + 0.9, C["muted"], 1); y += 0.9
        self.line(s, M + lw / 6, y, M + 5 * lw / 6, y, C["muted"], 1)
        aw = (lw - 0.2) / 3
        for i, (n, p) in enumerate([("Ervin Bahar P.", "Medium Vessel, Man Hauler, Bus Scania"), ("Tiberias K. Simu", "Small Suppeq, Truck Others"),
                                    ("Eurico", "Trailer, Overseas Zone")]):
            x = M + i * (aw + 0.1); me = i == 1
            self.line(s, x + aw / 2, y, x + aw / 2, y + 0.15, C["muted"], 1)
            if me:
                self.rect(s, x, y + 0.15, aw, 1.1, fill=C["orange_tint"], line=C["orange"], lw=1.5)
            else:
                self.card(s, x, y + 0.15, aw, 1.1, bar=False)
            self.text(s, x + 0.08, y + 0.25, aw - 0.16, 0.95, [[(n, {"bold": True, "size": 8.8})], [(p, {"size": 7.6, "color": C["orange_dk"] if me else C["text2"]})]],
                      color=C["orange_dk"] if me else C["black"], align=PP_ALIGN.CENTER)
        self.text(s, M, y + 1.4, lw, 0.5, "Associate", size=7.6, color=C["muted"], align=PP_ALIGN.CENTER)
        x = M + lw + 0.5; w = W - M - x
        self.text(s, x, 1.6, 4, 0.25, "MY PRODUCTS", size=8, bold=True, color=C["muted"])
        pw = (w - 0.2) / 2
        for i, (n, d) in enumerate([("Small Suppeq", "Water Truck (WT20, WT35), Lube Truck (LT08), Fuel Truck (FT20) and similar supporting units"),
                                    ("Truck Others", "Mining Crane, Fire Truck, Mixer Truck, Tyre Truck, Steaming Truck and other special-purpose trucks")]):
            xx = x + i * (pw + 0.2)
            self.card(s, xx, 1.95, pw, 1.35)
            self.text(s, xx + 0.25, 2.07, pw - 0.4, 0.3, n, size=11, bold=True, color=C["black"])
            self.text(s, xx + 0.25, 2.42, pw - 0.4, 0.8, d, size=8.2, color=C["text2"])
        self.evidence(s, x, 3.5, w, 3.35, "Product photos")

    def rr_b(self):
        s = self.new()
        self.title(s, "My Work Has Two Sides: Strategic Sets the Direction, Supply Chain Delivers the Order", "Roles & Responsibility · Strategic vs Supply Chain")
        y = 1.55; cw = (W - 2 * M - 0.2) / 2
        cats = [("A. STRATEGIC", C["orange"], "Decides direction — output feeds management decisions",
                 ["Pricing & Market Analysis", "Market Study & Program Development", "Market Program for Support Sales", "Performance Reporting (GP, PICA)", "Innovation & Development"]),
                ("B. SUPPLY CHAIN", C["navy"], "Runs the transaction — output keeps operations moving",
                 ["Quotation & Production Coordination", "Operational Reporting", "Internal & External Coordination", "Routine Meetings & Communication"])]
        for i, (n, col, d, items) in enumerate(cats):
            x = M + i * (cw + 0.2)
            self.rect(s, x, y, cw, 0.3, fill=col)
            self.text(s, x + 0.15, y, cw, 0.3, [[(n, {"bold": True}), ("   " + d, {"size": 7.6})]], size=9, color=C["white"], anchor=MSO_ANCHOR.MIDDLE)
            half = (len(items) + 1) // 2
            self.bullets(s, x + 0.1, y + 0.38, cw / 2 - 0.1, 0.8, items[:half], size=8, dot=col, spacing=1.5)
            self.bullets(s, x + cw / 2, y + 0.38, cw / 2 - 0.1, 0.8, items[half:], size=8, dot=col, spacing=1.5)
        # SIPOC
        y = 2.85
        self.text(s, M, y, 4, 0.22, "SIPOC", size=8, bold=True, color=C["muted"])
        hdr = ["Supplier", "Input", "Process", "Output", "Customer"]
        lab = 1.25; cw5 = (W - 2 * M - lab) / 5; y += 0.25
        rows = [("Supply Chain", C["gray_dk"], ["BC / Sales, UTE (Patria), non-Patria vendors, forwarders", "PCR demand, preliminary drawing, PO / PJB, QFD",
                 "Quote request → close won → PO Interco → production monitoring → billing", "Quotation, PO Interco, BAST / BAPB, invoice", "Sales team, end customer, Finance"]),
                ("Strategic", C["black"], ["CRM, sales & GP data, market sources, UTPE", "Market data, competitor specs, price history, GP results",
                 "Market study, price analysis, sales programs, GP & PICA reporting", "Market study, CRM price list, sales tools, GP / PICA report", "Management / BOD, Sales team"])]
        self.rect(s, M + lab + 2 * cw5, y, cw5, 0.28 + 2 * 0.72, fill=C["proc_bg"], line=C["proc_ln"], lw=1)
        for j, h in enumerate(hdr):
            self.text(s, M + lab + j * cw5, y, cw5, 0.28, h.upper(), size=7.8, bold=True, color=C["black"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for r, (n, col, cells) in enumerate(rows):
            yy = y + 0.28 + r * 0.72
            self.rect(s, M, yy + 0.04, lab - 0.1, 0.64, fill=col)
            self.text(s, M, yy + 0.04, lab - 0.1, 0.64, n, size=8.5, bold=True, color=C["white"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self.line(s, M + lab, yy, W - M, yy, C["line"], 0.5)
            for j, c in enumerate(cells):
                self.text(s, M + lab + j * cw5 + 0.08, yy + 0.06, cw5 - 0.16, 0.62, c, size=7.6, bold=(j == 2), color=C["text"], anchor=MSO_ANCHOR.MIDDLE)
        # bottom
        y = 5.15; bw3 = (W - 2 * M - 0.4) / 3
        boxes = [("Dimensions", ["Financial: non-Patria procurement budget, COGS per unit, target margin per PO", "Non-financial: vendors & customers handled, documents verified, report frequency, meetings & site visits"]),
                 ("Working Relationships", ["Internal: Sales & Marketing, Warehouse, Finance & Accounting, Legal, Procurement", "External: customers, Patria & non-Patria vendors, forwarders, UTPE"]),
                 ("Work Challenges", ["Keep pricing, procurement and report data accurate", "Align many functions on tight, shifting deadlines", "Adapt to changing market and stakeholder needs"])]
        for i, (n, items) in enumerate(boxes):
            x = M + i * (bw3 + 0.2)
            self.rect(s, x, y, bw3, 1.75, fill=C["pale"] if self.t["bg"] == "FFFFFF" else "FFFFFF")
            self.text(s, x + 0.18, y + 0.12, bw3 - 0.3, 0.25, n, size=9.5, bold=True, color=C["black"])
            self.bullets(s, x + 0.18, y + 0.45, bw3 - 0.3, 1.25, items, size=7.8, dot=C["gray_dk"])

    def step_cards(self, s, steps, color, y, h, ev_h, start=1, cols=None):
        n = len(steps); cols = cols or n
        cw = (W - 2 * M - (cols - 1) * 0.18) / cols
        for i, (t, items) in enumerate(steps):
            x = M + i * (cw + 0.18)
            self.card(s, x, y, cw, h, accent=color)
            self.badge_num(s, x + 0.2, y + 0.18, start + i, color)
            self.text(s, x + 0.62, y + 0.14, cw - 0.75, 0.42, t, size=9.5, bold=True, color=C["black"], anchor=MSO_ANCHOR.MIDDLE)
            self.bullets(s, x + 0.2, y + 0.65, cw - 0.35, h - ev_h - 0.8, items, size=7.8, dot=color, spacing=2)
            if ev_h:
                self.evidence(s, x + 0.2, y + h - ev_h - 0.18, cw - 0.38, ev_h)

    def jde_s1(self):
        s = self.new()
        self.title(s, "Strategy Starts With Data: I Frame the Question Before I Read the Market", "Job Description Execution · Strategic (1/2)")
        self.pill(s, W - M - 1.1, 1.0, "Steps 1–4", C["orange"], 1.1)
        self.step_cards(s, [
            ("Define Market & Objectives", ["Set the business question for Small Suppeq / Truck Others", "Agree scope and target segment with Dept Head"]),
            ("Design Market Study Plan", ["Choose data sources and method", "Set timeline and owners"]),
            ("Collect Internal & External Data", ["CRM, PCR, GP history, sales feedback", "Competitor specs and price points"]),
            ("Analyse Market & Competitors", ["Market size / share, segmentation, positioning", "Map competitor products against ours"]),
        ], C["orange"], 1.6, 5.3, 3.1)

    def jde_s2(self):
        s = self.new()
        self.title(s, "Analysis Only Counts When It Turns Into Tools, Reports and Decisions", "Job Description Execution · Strategic (2/2)")
        self.pill(s, W - M - 1.1, 1.0, "Steps 5–8", C["orange"], 1.1)
        self.step_cards(s, [
            ("Build Sales Tools & Programs", ["Pamphlets, incentive programs, campaigns", "Standard price list in CRM, event support"]),
            ("Prepare Performance Reports", ["GP Report and Performance Report", "PICA analysis on gaps"]),
            ("Interpret Results & Recommend", ["Turn findings into pricing / program advice", "Present to management"]),
            ("Innovation & Development", ["Project Charter for innovation projects", "Training (Patria Mover), site visit reports"]),
        ], C["orange"], 1.6, 3.55, 1.55, start=5)
        y = 5.35; hdr = ["Deliverables", "Departments Involved", "Challenges", "Improvements"]
        cells = ["Market study, sales tools, CRM price list, GP & PICA reports", "Sales, Finance & Accounting, UTPE, Marketing Communication",
                 "Scattered data sources; limited competitor price visibility", "Single dashboard for GP and pipeline; regular competitor mapping"]
        cw = (W - 2 * M) / 4
        self.rect(s, M, y, W - 2 * M, 0.3, fill=C["black"])
        for j in range(4):
            self.text(s, M + j * cw + 0.15, y, cw - 0.2, 0.3, hdr[j].upper(), size=7.8, bold=True, color=C["white"], anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, M + j * cw + 0.15, y + 0.38, cw - 0.3, 0.9, cells[j], size=8, color=C["text"])
            if j:
                self.line(s, M + j * cw, y + 0.35, M + j * cw, y + 1.4, C["line"], 0.5)
        self.line(s, M, y + 1.45, W - M, y + 1.45, C["line"], 0.75)

    def jde_c1(self):
        s = self.new()
        self.title(s, "Before Any Price Goes Out, I Clarify the Demand and Read the Drawing", "Job Description Execution · Supply Chain (1/2)")
        self.pill(s, W - M - 2.1, 1.0, "Main Process · Steps 1–3", C["gray_dk"], 2.1)
        self.step_cards(s, [
            ("Demand Check", ["See demand rise in PCR", "Clarify with BC: reason to buy (expansion / new business / replacement)",
                              "Project location; fast or low-cost priority", "Expected price and estimated PO date"]),
            ("Preliminary Drawing Review", ["Check unit function and how it works", "Flag critical main components on price and lead time",
                                            "Consult Application Engineering"]),
            ("Request Quote to UTE", ["Attach prelim, expected price & lead time, end customer, expected delivery",
                                      "Check standard GP (13.6% — pre-negotiation price)", "Track SLA, escalate if late; log result in PCR"]),
        ], C["gray_dk"], 1.6, 5.3, 2.55)

    def jde_c2(self):
        s = self.new()
        self.title(s, "Winning the Order Is Half the Job — Monitoring Keeps Delivery on Schedule", "Job Description Execution · Supply Chain (2/2)")
        self.pill(s, W - M - 2.1, 1.0, "Main Process · Steps 4–5", C["gray_dk"], 2.1)
        self.step_cards(s, [
            ("Close Won → CPO", ["Check PO / prelim / PJB in CRM; request QFD from UTE; collect special requests & painting style from BC",
                                 "RFD from QFD becomes PO Interco due date",
                                 "Close won: Sales Mgr → Mkt Associate → Mkt Mgr → GM Mkt · PO Interco approval: Mkt Mgr → GM Mkt → Mkt Director"]),
            ("Monitoring", ["Track production and update sales (factory visit or photos from Nandu); match QFD vs actual",
                            "Escalate when RFD is at risk; request painting style ≥1 month before RFD", "Accompany BC at customer FAT"]),
        ], C["gray_dk"], 1.45, 2.2, 0, start=4)
        y = 3.82
        self.text(s, M, y, 4, 0.22, "INTERNAL MONITORING", size=8, bold=True, color=C["muted"])
        mons = [("Pipeline–OSPO", "Dashboard, updated monthly"), ("CRM Staging Pipeline", "No prelim = early talks; prelim = validate & quote"),
                ("Quote Price to UTE", "Detail, average SLA, initial price"), ("Progress Billing", "Check staging, escalate")]
        mw = (W - 2 * M - 3 * 0.15) / 4
        for i, (n, d) in enumerate(mons):
            x = M + i * (mw + 0.15)
            self.card(s, x, y + 0.27, mw, 0.75, accent=C["navy"])
            self.text(s, x + 0.17, y + 0.33, mw - 0.25, 0.25, n, size=8.6, bold=True, color=C["black"])
            self.text(s, x + 0.17, y + 0.58, mw - 0.25, 0.42, d, size=7.6, color=C["text2"])
        y = 5.05
        self.text(s, M, y, 5, 0.22, "BILLING PROCESS CHAIN", size=8, bold=True, color=C["muted"])
        chains = [("At UTE", C["gray_dk"], ["PO Interco", "PRO", "PR", "PO", "Drawing", "BOM", "LPPB", "PB", "Fabrication", "Assembly", "QC passed", "GR Teco", "GI"]),
                  ("At Triatra", C["orange"], ["GR", "DO / TO", "GI", "BAST", "BAST approval", "Invoice"])]
        for r, (lab, col, steps) in enumerate(chains):
            yy = y + 0.3 + r * 0.7
            self.rect(s, M, yy, 1.0, 0.45, fill=col)
            self.text(s, M, yy, 1.0, 0.45, lab, size=8, bold=True, color=C["white"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x0 = M + 1.1; avail = W - M - x0; sw = (avail - (len(steps) - 1) * 0.12) / len(steps) if r == 0 else (avail - 5 * 0.12) / 6 * 0.62
            for k, st in enumerate(steps):
                x = x0 + k * (sw + 0.12)
                self.rect(s, x, yy, sw, 0.45, fill="FFFFFF", line=col, lw=0.75)
                self.text(s, x, yy, sw, 0.45, st, size=7, bold=True, color=C["text"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                if k < len(steps) - 1:
                    self.text(s, x + sw - 0.02, yy, 0.16, 0.45, "›", size=9, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def build(self):
        self.agenda()
        self.divider(1, "Company Profile"); self.company()
        self.divider(2, "Site Operation"); self.site()
        self.divider(3, "Organization Chart"); self.org()
        subs = ["A — Position & Products", "B — Strategic vs Supply Chain"]
        self.divider(4, "Roles & Responsibility", subs, 0); self.rr_a(); self.rr_b()
        self.divider(5, "Job Description Execution", ["Strategic", "Supply Chain"], 0)
        self.jde_s1(); self.jde_s2(); self.jde_c1(); self.jde_c2()
        return self.prs


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    for key, th in THEMES.items():
        Deck(th).build().save(f"{out}/NEDP_Deck_{key}.pptx")
        print("saved", key)
