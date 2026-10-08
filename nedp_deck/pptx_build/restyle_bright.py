"""Restyle the user's edited V1 deck into the bright Ref5 theme, keeping every text/shape edit."""
import json, re, io, sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
SRC, OUT = sys.argv[1], sys.argv[2]
H = open("/home/user/6month/nedp_deck/NEDP_Deck_Ref5_Bright.html").read()
CLS = re.findall(r'<section[^>]*class="slide([^"]*)"', H)
items = json.load(open("/tmp/claude-0/pp5/items.json"))
htxt = [set(" ".join(r["t"] for it in s for r in it["runs"]).lower().split()) for s in items]

TXT = {"F4F1EE": "1B1B1B", "F0ECE8": "1B1B1B", "EFE9E4": "1B1B1B", "A9A29C": "5E5A55", "A79F98": "5E5A55",
       "6F6963": "9A948D", "FFD3C0": "FF5A1F", "FFB089": "FF7A45", "F5C26B": "E08A00", "F2C14E": "E08A00"}
def lum(h): r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4)); return (0.299*r + 0.587*g + 0.114*b) / 255
def fillmap(h):
    if lum(h) > 0.3: return h
    r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4))
    return "FFF1EA" if r - b > 25 else "FBF7F2"

p = Presentation(SRC)
W, Hh = p.slide_width, p.slide_height
white = io.BytesIO(); Image.new("RGB", (1920, 1080), "white").save(white, "PNG")
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

def recolor_xml(el, divider=False):
    for c in el.iter(NS + "srgbClr"):
        v = c.get("val").upper(); par = c.getparent()
        tag = par.tag.split("}")[1]
        in_text = any(a.tag.split("}")[1] in ("rPr", "defRPr", "endParaRPr") for a in [par] + list(par.iterancestors())[:2])
        if divider and in_text: c.set("val", "FFFFFF"); continue
        if in_text: c.set("val", TXT.get(v, v))
        elif tag == "solidFill" or tag == "ln" or par.getparent() is not None:
            nv = fillmap(v)
            if any(a.tag.endswith("}ln") for a in par.iterancestors()) and lum(v) < 0.3: nv = "ECE6DE"
            c.set("val", nv)

for i, s in enumerate(p.slides):
    pics = [sh for sh in s.shapes if sh.shape_type == 13 and sh.left <= 0 and sh.top <= 0 and sh.width >= W * 0.95]
    words = set(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).lower().split())
    k = max(range(len(htxt)), key=lambda j: len(words & htxt[j]) / (len(words | htxt[j]) or 1))
    score = len(words & htxt[k]) / (len(words | htxt[k]) or 1)
    native = s.slide_layout.name != "Blank" or score < 0.5
    divider = (not native) and "div" in CLS[k]
    if pics and not native:
        im = Image.open(f"/tmp/claude-0/pp5/bg{k}.png").convert("RGB"); b = io.BytesIO(); im.save(b, "JPEG", quality=88)
        pics[0]._element.blipFill  # ensure pic
        rId = pics[0]._element.blipFill.blip.rEmbed
        part = s.part.related_part(rId); part._blob = b.getvalue()
    elif pics:
        part = s.part.related_part(pics[0]._element.blipFill.blip.rEmbed); part._blob = white.getvalue()
    # slide background -> white
    bg = s.background.fill; bg.solid(); 
    from pptx.dml.color import RGBColor; bg.fore_color.rgb = RGBColor(255, 255, 255)
    recolor_xml(s.shapes._spTree, divider)
    # tables: orange header, zebra body
    for sh in s.shapes:
        if sh.has_table:
            for ri, row in enumerate(sh.table.rows):
                for cell in row.cells:
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string("FF5A1F" if ri == 0 else ("FFFFFF" if ri % 2 else "FBF7F2"))
                    for para in cell.text_frame.paragraphs:
                        for r in para.runs:
                            cur = None
                            try: cur = str(r.font.color.rgb)
                            except Exception: pass
                            if ri == 0: r.font.color.rgb = RGBColor(255, 255, 255)
                            elif cur in (None, "FFFFFF", "1B1B1B", "5E5A55", "9A948D"): r.font.color.rgb = RGBColor.from_string("1B1B1B" if cur in (None, "FFFFFF", "1B1B1B") else cur)
        if sh.has_chart:
            cp = sh.chart.part
            x = cp._element
            recolor_xml(x)
    print(i + 1, "native" if native else f"html#{k+1}{' div' if divider else ''}", round(score, 2))
# chart xml text colours
for part in p.part.package.iter_parts():
    if part.partname.startswith("/ppt/charts/") and part.partname.endswith(".xml") and hasattr(part, "_element"):
        recolor_xml(part._element)
p.save(OUT)
