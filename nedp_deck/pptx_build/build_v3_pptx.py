"""Editable PPTX from v3/Deck_v3.html. Run node pptx_build/ex7.js first (writes /tmp/claude-0/pp7)."""
import base64, io, json, re, sys
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Emu, Pt
PP = '/tmp/claude-0/pp7/'
ITEMS = json.load(open(PP + 'items.json')); IMGS = json.load(open(PP + 'imgs.json'))
K = 12192000 / 1280
def E(v): return Emu(round(v * K))
def img(d):
    im = Image.open(io.BytesIO(base64.b64decode(d['src'].split(',', 1)[1])))
    im = im.convert('RGBA') if im.mode in ('RGBA', 'LA', 'P') else im.convert('RGB')
    W, H = im.size; c = d['crop']
    im = im.crop((max(0, round(c[0]*W)), max(0, round(c[1]*H)), min(W, round(c[2]*W)), min(H, round(c[3]*H))))
    o = io.BytesIO(); im.save(o, 'PNG') if im.mode == 'RGBA' else im.save(o, 'JPEG', quality=92); o.seek(0); return o
P = Presentation(); P.slide_width, P.slide_height = 12192000, 6858000
for k, (items, imgs) in enumerate(zip(ITEMS, IMGS)):
    s = P.slides.add_slide(P.slide_layouts[6])
    o = io.BytesIO(); Image.open(f'{PP}bg{k}.png').convert('RGB').save(o, 'JPEG', quality=90); o.seek(0)
    s.shapes.add_picture(o, 0, 0, P.slide_width, P.slide_height)
    for d in imgs: s.shapes.add_picture(img(d), E(d['x']), E(d['y']), E(d['w']), E(d['h']))
    for it in items:
        tb = s.shapes.add_textbox(E(it['x'] - 1), E(it['y'] - 1), E(it['w'] * 1.06 + 4), E(it['h'] + 2))
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = MSO_ANCHOR.TOP
        for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, a, 0)
        p = tf.paragraphs[0]
        for r in it['runs']:
            run = p.add_run(); run.text = r['t'].upper() if r['up'] else r['t']; f = run.font
            f.size = Pt(r['fs'] * .75); f.bold = r['b']; f.italic = r.get('it', False)
            f.name = 'Georgia' if 'Georgia' in r.get('f', '') else 'Arial'
            f.color.rgb = RGBColor(*[int(float(x)) for x in re.findall(r'[\d.]+', r['c'])[:3]])
P.save(sys.argv[1]); print('saved', len(P.slides))
