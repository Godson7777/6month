"""Bright deck from the user's V1: every HTML-designed slide is rebuilt from NEDP_Deck_Ref5_Bright.html
(sharp background + native pictures + editable Arial text), native F.lli Ferrari slides are recoloured,
switched to Arial and their appendix grids enlarged; footer and page numbers are unified.
Run after:  node pptx_build/ex6.js   (writes /tmp/claude-0/pp5/{bg*.png,items.json,imgs.json})
Usage:      python3 pptx_build/build_bright_final.py V1.pptx OUT.pptx"""
import base64, copy, io, json, re, sys
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
PP = "/tmp/claude-0/pp5/"
H = open("/home/user/6month/nedp_deck/NEDP_Deck_Ref5_Bright.html").read()
CLS = re.findall(r'<section[^>]*class="slide([^"]*)"', H)
ITEMS = json.load(open(PP + "items.json"))
IMGS = json.load(open(PP + "imgs.json"))
K = 12192000 / 1280                       # EMU per HTML px
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FONT = "Arial"
ORANGE, INK, INK2, MUTED = "FF5A1F", "1B1B1B", "5E5A55", "9A948D"


def words(items):
    return set(" ".join(r["t"] for it in items for r in it["runs"]).lower().split())
HW = [words(s) for s in ITEMS]


def colr(c):
    return RGBColor(*[int(float(x)) for x in re.findall(r"[\d.]+", c)[:3]])


def img_stream(d):
    im = Image.open(io.BytesIO(base64.b64decode(d["src"].split(",", 1)[1])))
    im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
    W0, H0 = im.size; c = d["crop"]
    im = im.crop((max(0, round(c[0] * W0)), max(0, round(c[1] * H0)), min(W0, round(c[2] * W0)), min(H0, round(c[3] * H0))))
    if d["inv"]:
        a = im.convert("RGBA").getchannel("A"); im = Image.new("RGBA", im.size, (255, 255, 255, 255)); im.putalpha(a)
    o = io.BytesIO()
    im.save(o, "PNG", optimize=True) if im.mode == "RGBA" else im.save(o, "JPEG", quality=93)
    o.seek(0); return o


def bg_stream(k):
    o = io.BytesIO(); Image.open(f"{PP}bg{k}.png").convert("RGB").save(o, "JPEG", quality=92); o.seek(0); return o


PAGE = re.compile(r"^\s*\d{1,2}\s*/\s*\d{1,2}\s*$")


def build_html_slide(s, k, page, total):
    """Replace everything on slide s with HTML slide k."""
    tree = s.shapes._spTree
    RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for el in list(tree)[2:]:
        rids = {a for e in el.iter() for n, a in e.attrib.items() if n.startswith(RID)}
        tree.remove(el)
        for rid in rids:
            try: s.part.drop_rel(rid)
            except Exception: pass
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    s.shapes.add_picture(bg_stream(k), 0, 0, Emu(12192000), Emu(6858000))
    for d in IMGS[k]:
        s.shapes.add_picture(img_stream(d), Emu(int(d["x"] * K)), Emu(int(d["y"] * K)), Emu(int(d["w"] * K)), Emu(int(d["h"] * K)))
    for t in ITEMS[k]:
        runs = [dict(r) for r in t["runs"]]
        if PAGE.match("".join(r["t"] for r in runs)):
            runs = [dict(runs[0], t=f"{page:02d} "), dict(runs[-1], t=f"/ {total}")]
        tb = s.shapes.add_textbox(Emu(int(t["x"] * K)), Emu(int((t["y"] - t["h"] * 0.08) * K)),
                                  Emu(int((t["w"] * 1.06 + 6) * K)), Emu(int(t["h"] * K)))
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
        para = tf.paragraphs[0]; para.alignment = PP_ALIGN.LEFT
        for j, x in enumerate(runs):
            r = para.add_run(); txt = x["t"].upper() if x.get("up") else x["t"]
            r.text = txt.lstrip() if j == 0 else txt
            f = r.font; f.size = Pt(x["fs"] * 0.75 * 0.97); f.bold = x["b"]; f.name = FONT; f.color.rgb = colr(x["c"])


# ---------------- native-slide restyle ----------------
TXT = {"F4F1EE": INK, "F0ECE8": INK, "EFE9E4": INK, "A9A29C": INK2, "A79F98": INK2,
       "6F6963": MUTED, "FFD3C0": ORANGE, "FFB089": "FF7A45", "F5C26B": "E08A00", "F2C14E": "E08A00"}


def lum(h):
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4)); return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def fillmap(h):
    if lum(h) > 0.3: return h
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return "FFF1EA" if r - b > 25 else "FBF7F2"


TEXT_TAGS = ("rPr", "defRPr", "endParaRPr")


def recolor(root):
    for c in root.iter(NS + "srgbClr"):
        v = c.get("val").upper(); par = c.getparent()
        if any(a.tag.split("}")[1] in TEXT_TAGS for a in [par] + list(par.iterancestors())[:2]):
            c.set("val", TXT.get(v, v))
        else:
            nv = fillmap(v)
            if any(a.tag.endswith("}ln") for a in par.iterancestors()) and lum(v) < 0.3: nv = "ECE6DE"
            c.set("val", nv)


def arial(root):
    """Every run, paragraph default and end-of-paragraph property -> Arial."""
    for tag in TEXT_TAGS:
        for rp in root.iter(NS + tag):
            for sub in ("latin", "ea", "cs"):
                el = rp.find(NS + sub)
                if el is None:
                    el = etree.SubElement(rp, NS + sub)
                    # schema order: fill/effects first, then latin/ea/cs, then hyperlinks
                    hl = [x for x in rp if x.tag in (NS + "hlinkClick", NS + "hlinkMouseOver", NS + "rtl", NS + "extLst")]
                    if hl: hl[0].addprevious(el)
                el.set("typeface", FONT)
                for a in ("panose", "pitchFamily", "charset"):
                    el.attrib.pop(a, None)


def no_spautofit(root):
    for bp in root.iter(NS + "bodyPr"):
        for af in list(bp):
            if af.tag == NS + "spAutoFit": bp.remove(af)


def restyle_native(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    W = 12192000
    for sh in s.shapes:   # full-bleed dark picture backgrounds -> white
        if sh.shape_type == 13 and sh.left <= 0 and sh.top <= 0 and sh.width >= W * 0.95:
            o = io.BytesIO(); Image.new("RGB", (64, 36), "white").save(o, "PNG")
            s.part.related_part(sh._element.blipFill.blip.rEmbed)._blob = o.getvalue()
    recolor(s.shapes._spTree); arial(s.shapes._spTree); no_spautofit(s.shapes._spTree)
    for sh in s.shapes:
        if sh.has_table:
            for ri, row in enumerate(sh.table.rows):
                for cell in row.cells:
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(ORANGE if ri == 0 else ("FFFFFF" if ri % 2 else "FBF7F2"))
                    for para in cell.text_frame.paragraphs:
                        for r in para.runs:
                            try: cur = str(r.font.color.rgb)
                            except Exception: cur = None
                            if ri == 0: r.font.color.rgb = RGBColor(255, 255, 255)
                            elif cur in (None, "FFFFFF"): r.font.color.rgb = RGBColor.from_string(INK)
                            r.font.name = FONT
        if sh.has_chart:
            recolor(sh.chart.part._element); arial(sh.chart.part._element)


def grow_appendix_grid(s, top=93.6, bottom=492.0, gap=9.0):
    """Stretch the 3-row photo grid of an appendix slide to fill the free height (crop keeps photos undistorted)."""
    P = 12700
    cards = sorted({round(sh.top / P, 1) for sh in s.shapes if sh.shape_type == 1 and 90 < sh.height / P < 120})
    if len(cards) < 2: return
    rows = len(cards); ch = (bottom - top - gap * (rows - 1)) / rows
    for sh in list(s.shapes):
        t = sh.top / P
        r = next((i for i, c in enumerate(cards) if c - 0.5 <= t <= c + 105), None)
        if r is None or t < 90 or t > 480: continue
        off = t - cards[r]; nt = top + r * (ch + gap)
        if sh.shape_type == 1 and 90 < sh.height / P < 120:            # card
            sh.top = Emu(int(nt * P)); sh.height = Emu(int(ch * P))
        elif sh.shape_type == 13:                                        # photo
            old_h = sh.height / P; new_h = ch - (105.8 - old_h)
            pic = sh._element; blip = pic.blipFill.blip
            iw, ih = Image.open(io.BytesIO(s.part.related_part(blip.rEmbed).blob)).size
            cl, cr, ct, cb = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
            vis_w, vis_h = iw * (1 - cl - cr), ih * (1 - ct - cb)
            target = (sh.width / P) / new_h
            if vis_w / vis_h > target:     # too wide -> crop sides
                fw = vis_h * target / iw; extra = (1 - cl - cr) - fw
                sh.crop_left, sh.crop_right = cl + extra / 2, cr + extra / 2
            else:                           # too tall -> crop top/bottom
                fh = vis_w / target / ih; extra = (1 - ct - cb) - fh
                sh.crop_top, sh.crop_bottom = ct + extra / 2, cb + extra / 2
            sh.top = Emu(int(nt * P)); sh.height = Emu(int(new_h * P))
        else:                                                            # label / tag below the photo
            sh.top = Emu(int((nt + off + (ch - 105.8)) * P)) if off > 60 else Emu(int((nt + off) * P))


def footer_native(s, page, total, ref_items):
    """Page number 'NN / total' and the deck footer at the same place and style as the HTML slides."""
    P = 12700
    for sh in list(s.shapes):
        if not sh.has_text_frame or sh.top / P < 470: continue
        tx = sh.text_frame.text.strip()
        if re.fullmatch(r"\d{1,2}(\s*/\s*\d{1,2})?", tx) or (tx and tx == tx.upper() and len(tx) < 70 and sh.top / P > 505):
            sh._element.getparent().remove(sh._element)       # old page number / old upper-case footer label
    def band_free(x0, y0, x1, y1):
        return not any(sh.left / P < x1 and (sh.left + sh.width) / P > x0 and sh.top / P < y1 and (sh.top + sh.height) / P > y0
                       for sh in s.shapes)
    for t in ref_items:
        txt = "".join(r["t"] for r in t["runs"])
        if not (PAGE.match(txt) or txt.strip().upper() == "NEDP MID YEAR REVIEW 2026"): continue
        x0, y0 = t["x"] * 0.75, t["y"] * 0.75
        if not PAGE.match(txt) and not band_free(x0, y0 - 2, x0 + t["w"] * 0.75, y0 + t["h"] * 0.75): continue
        runs = t["runs"] if not PAGE.match(txt) else [dict(t["runs"][0], t=f"{page:02d} "), dict(t["runs"][-1], t=f"/ {total}")]
        tb = s.shapes.add_textbox(Emu(int(t["x"] * K)), Emu(int((t["y"] - t["h"] * 0.08) * K)),
                                  Emu(int((t["w"] * 1.06 + 6) * K)), Emu(int(t["h"] * K)))
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]; para.alignment = PP_ALIGN.LEFT
        for j, x in enumerate(runs):
            r = para.add_run(); tx = x["t"].upper() if x.get("up") else x["t"]; r.text = tx.lstrip() if j == 0 else tx
            r.font.size = Pt(x["fs"] * 0.75 * 0.97); r.font.bold = x["b"]; r.font.name = FONT; r.font.color.rgb = colr(x["c"])


def theme_arial(prs):
    for part in prs.part.package.iter_parts():
        if part.partname.startswith("/ppt/theme/") and part.partname.endswith(".xml"):
            x = part.blob.decode("utf8")
            x = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*"', r'\1Arial"', x)
            part._blob = x.encode("utf8")


# ---------------- assemble ----------------
p = Presentation(SRC)
plan = []
for i, s in enumerate(p.slides):
    w = set(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).lower().split())
    k = max(range(len(HW)), key=lambda j: len(w & HW[j]) / (len(w | HW[j]) or 1))
    sc = len(w & HW[k]) / (len(w | HW[k]) or 1)
    plan.append(k if (s.slide_layout.name == "Blank" and sc >= 0.5) else None)

# insert About Me (html 3) and Onboarding Journey (html 4) after the Personal Information divider
pos = next(i for i, k in enumerate(plan) if k is not None and "PERSONAL" in " ".join(r["t"] for it in ITEMS[k] for r in it["runs"]).upper() and "div" in CLS[k]) + 1
blank_layout = p.slides[pos - 1].slide_layout
for k_new in (3, 2):
    ns = p.slides.add_slide(blank_layout)
    for ph in list(ns.placeholders): ph._element.getparent().remove(ph._element)
    lst = p.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    plan.insert(pos, k_new)

total = len(plan)
ref_footer = ITEMS[2]
for i, (s, k) in enumerate(zip(p.slides, plan)):
    if k is not None:
        build_html_slide(s, k, i + 1, total)
    else:
        restyle_native(s)
        title = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).upper()
        if "APPENDIX" in title and "MODEL PHOTOS" in title:
            grow_appendix_grid(s)
        footer_native(s, i + 1, total, ref_footer)
    print(i + 1, f"html#{k + 1}" if k is not None else "native")
theme_arial(p)
p.save(OUT)
print("saved", OUT, total, "slides")
