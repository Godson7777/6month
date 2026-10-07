import json,re
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
D=json.load(open('items.json'));P=Presentation();P.slide_width=Emu(12192000);P.slide_height=Emu(6858000)
k=12192000/1280
def col(c):
    m=re.findall(r'[\d.]+',c);return RGBColor(*[int(float(x)) for x in m[:3]])
for i,items in enumerate(D):
    s=P.slides.add_slide(P.slide_layouts[6])
    s.shapes.add_picture(f'bg{i}.jpg',0,0,P.slide_width,P.slide_height)
    for t in items:
        tb=s.shapes.add_textbox(Emu(int(t['x']*k)),Emu(int(t['y']*k)),Emu(int((t['w']*1.06+8)*k)),Emu(int(t['h']*k)))
        tf=tb.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
        lines=[[]]
        for r0 in t.get('runs') or [{'t':t['t'],'fs':t['fs'],'b':t['fw'],'c':t['col']}]:
            parts=r0['t'].split('\n')
            for pi,pt in enumerate(parts):
                if pi>0: lines.append([])
                if pt: lines[-1].append(dict(r0,t=pt.upper() if r0.get('up') else pt))
        lines=[l for l in lines if any(x['t'].strip() for x in l)] or [[]]
        for j,line in enumerate(lines):
            para=tf.paragraphs[0] if j==0 else tf.add_paragraph()
            para.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(t['al'],PP_ALIGN.LEFT)
            for x in line:
                r=para.add_run();r.text=x['t'].lstrip() if line.index(x)==0 else x['t'];f=r.font;f.size=Pt(x['fs']*0.75*0.96);f.bold=x['b'];f.name='Arial';f.color.rgb=col(x['c'])
P.save('/home/user/6month/nedp_deck/NEDP_MidYear_Review_2026_Editable.pptx')
