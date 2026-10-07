import json,re
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
D=json.load(open('items.json'));P=Presentation();P.slide_width=Emu(12192000);P.slide_height=Emu(6858000)
k=12192000/1280
def col(c):
    m=re.findall(r'[\d.]+',c);return RGBColor(*[int(float(x)) for x in m[:3]])
for i,items in enumerate(D):
    s=P.slides.add_slide(P.slide_layouts[6])
    s.shapes.add_picture(f'bg{i}.png',0,0,P.slide_width,P.slide_height)
    for t in items:
        tb=s.shapes.add_textbox(Emu(int(t['x']*k)),Emu(int(t['y']*k)),Emu(int((t['w']+6)*k)),Emu(int(t['h']*k)))
        tf=tb.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
        for j,line in enumerate(t['t'].split('\n')):
            para=tf.paragraphs[0] if j==0 else tf.add_paragraph()
            para.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(t['al'],PP_ALIGN.LEFT)
            r=para.add_run();r.text=line;f=r.font;f.size=Pt(t['fs']*0.75*1.0);f.bold=t['fw'];f.name='Arial';f.color.rgb=col(t['col'])
P.save('/home/user/6month/nedp_deck/NEDP_MidYear_Review_2026.pptx')
