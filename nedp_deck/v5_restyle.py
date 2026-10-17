"""Apply 'Shaping the future' template look to 6M_V5.pptx without changing text:
template font everywhere; old NEDP chrome (logo, top/bottom rules, footer, kicker) replaced by
the template date + slide-number placeholders; headlines restyled like the user's own titles."""
import copy, sys, re, zipfile, shutil
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
FONT = 'Neue Haas Grotesk Text Pro'
SRC, OUT = sys.argv[1], sys.argv[2]
P = Presentation(SRC)
px = lambda v: round(v / 9525)
ref = P.slides[8]  # user's own slide with template date / number placeholders
DT = [s for s in ref.shapes if s.is_placeholder and s.placeholder_format.type in (16, 13)]
OLD = re.compile(r'^(NEDP 6-MONTH REVIEW|NEDP MID YEAR REVIEW 2026|\d\d / \d\d$)')
for idx, s in enumerate(P.slides, 1):
    shapes = list(s.shapes)
    has_old = any(sh.has_text_frame and OLD.match(sh.text_frame.text.strip()) for sh in shapes)
    if not has_old: continue
    head = None
    for sh in shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        t = sh.text_frame.text.strip() if sh.has_text_frame else ''
        kill = (sh.shape_type == 13 and y < 40 and x < 120) or (sh.shape_type == 9 and h == 0 and y in (56, 692)) \
            or (t and OLD.match(t)) or (sh.has_text_frame and 70 <= y <= 78 and h <= 16 and t.isupper())
        if kill: sh._element.getparent().remove(sh._element); continue
        if sh.has_text_frame and 86 <= y <= 94 and h >= 36 and w > 1000: head = sh
    for d in DT: s.shapes._spTree.append(copy.deepcopy(d._element))
    if head is not None:
        head.left, head.top, head.width, head.height = Emu(36 * 9525), Emu(58 * 9525), Emu(1207 * 9525), Emu(62 * 9525)
        for p in head.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(26); r.font.bold = True; r.font.italic = False; r.font.color.rgb = RGBColor(0, 0, 0)
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(36 * 9525), Emu(124 * 9525), Emu(1243 * 9525), Emu(124 * 9525))
        st = ln._element.find(qn('p:style'))
        if st is not None: ln._element.remove(st)
        ln.line.color.rgb = RGBColor(0, 0, 0); ln.line.width = Pt(1.5)
        for sh in s.shapes:  # lede line under old headline
            if sh.has_text_frame and px(sh.top) == 132: sh.top = Emu(130 * 9525)
P.save(OUT)
# font: every explicit Arial / Georgia -> template font (slides and charts)
tmp = OUT + '.tmp'; shutil.move(OUT, tmp)
with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        b = zi.read(it.filename)
        if it.filename.endswith('.xml') and (it.filename.startswith('ppt/slides/') or it.filename.startswith('ppt/charts/')):
            b = re.sub(rb'typeface="(Arial|Georgia|Liberation Sans)"', ('typeface="%s"' % FONT).encode(), b)
        zo.writestr(it, b)
import os; os.remove(tmp); print('ok')
