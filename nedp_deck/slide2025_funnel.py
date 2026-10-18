"""2025 funnel slide: market by class (Light/Small excluded) -> Medium brands treemap -> Heavy brands treemap."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree
FONT = 'Neue Haas Grotesk Text Pro'; E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string
ACC, GREY, LG = '1E9E57', '6B6B6B', 'BFBFBF'
FX = 16000  # assumed IDR per USD
CLS = ['Light', 'Small', 'Medium', 'Heavy']
U = {'Light': 53, 'Small': 75, 'Medium': 60, 'Heavy': 39}                       # database 2025
V = {'Light': 551005, 'Small': 1777124, 'Medium': 1759387, 'Heavy': 2951742}    # USD 2025
COL = {'Light': 'E3E3E3', 'Small': 'C9C9C9', 'Medium': '1E9E57', 'Heavy': '0B4D2A'}
rp = lambda usd: f'Rp {usd * FX / 1e9:,.1f} Bio'.replace(',', '#').replace('.', ',').replace('#', '.')
TU = sum(U.values()); TV = sum(V.values()); MHU = U['Medium'] + U['Heavy']; MHV = V['Medium'] + V['Heavy']
# brand units 2025 (top 3 from earlier data cut; Others = database total minus top 3)
MED = [('Sany Palfinger', 30), ('XCMG', 14), ('Palfinger', 5)]; MED.append(('Others', U['Medium'] - sum(b[1] for b in MED)))
HEV = [('Palfinger', 22), ('Amco Veba', 7), ('XCMG', 4)]; HEV.append(('Others', U['Heavy'] - sum(b[1] for b in HEV)))
SHADE = ['0B4D2A', '1E9E57', '6FCF97', 'C7EFD5']

P = Presentation(); P.slide_width, P.slide_height = 12192000, 6858000
s = P.slides.add_slide(P.slide_layouts[6])
def tx(x, y, w, h, parts, sz=12, al=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    if isinstance(parts, str): parts = [[(parts, False, '000000')]]
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = al
        for t, bold, c in para:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(c)
    return b
def line(x1, y1, x2, y2, c='000000', w=1.0):
    l = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)); l.line.color.rgb = rgb(c); l.line.width = Pt(w); return l
def rect(x, y, w, h, fill, line_c='FFFFFF', lw=1.5):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if fill: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else: r.fill.background()
    r.line.color.rgb = rgb(line_c); r.line.width = Pt(lw); return r
def tile(x, y, w, h, fill, top, bottom, dark):
    r = rect(x, y, w, h, fill); tf = r.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = E(6); tf.margin_top = tf.margin_bottom = E(5); tf.vertical_anchor = MSO_ANCHOR.TOP
    c = 'FFFFFF' if dark else '000000'
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; rr = p.add_run(); rr.text = top
    rr.font.size = Pt(13); rr.font.bold = True; rr.font.name = FONT; rr.font.color.rgb = rgb(c)
    # bottom label as separate textbox anchored at bottom-left
    tx(x + 5, y + h - 40, w - 10, 36, [[(bottom, True, c)]], 10.5 if w > 70 else 8.5, anchor=MSO_ANCHOR.BOTTOM)
def treemap(x, y, w, h, brands):
    T = sum(b[1] for b in brands); sh = [b[1] / T for b in brands]
    w0 = w * sh[0]; tile(x, y, w0, h, SHADE[0], f'~{sh[0]:.0%}', f'{brands[0][0]}, {brands[0][1]}', True)
    rest = sum(sh[1:]); rw = w - w0
    h1 = h * sh[1] / rest; tile(x + w0, y, rw, h1, SHADE[1], f'~{sh[1]:.0%}', f'{brands[1][0]}, {brands[1][1]}', True)
    w2 = rw * sh[2] / (sh[2] + sh[3])
    tile(x + w0, y + h1, w2, h - h1, SHADE[2], f'~{sh[2]:.0%}', f'{brands[2][0]}, {brands[2][1]}', False)
    tile(x + w0 + w2, y + h1, rw - w2, h - h1, SHADE[3], f'~{sh[3]:.0%}', f'{brands[3][0]}, {brands[3][1]}', False)
def banner(x, w, text, fill):
    b = rect(x, 214, w, 26, fill, fill, 1); tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = text
    r.font.size = Pt(12); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb('FFFFFF')

# ---------- header ----------
tx(36, 22, 200, 16, '10/18/2026', 11); tx(440, 22, 400, 16, 'Shaping the future of our industry', 11, PP_ALIGN.CENTER); tx(1144, 22, 100, 16, '13', 11, PP_ALIGN.RIGHT)
tx(36, 52, 900, 60, [[('Indonesian Import Folding Crane 2025', True, '000000')]], 34, anchor=MSO_ANCHOR.MIDDLE)
tx(940, 54, 304, 56, [[('Where is the value?', True, '000000')], [('And who holds it?', True, '000000')]], 16, PP_ALIGN.RIGHT)
line(28, 116, 1252, 116, '000000', 3)
tx(36, 122, 1208, 28, [[('Medium and Heavy are 44% of units but 67% of value, and each class has one clear leader', True, '000000')]], 19, anchor=MSO_ANCHOR.MIDDLE)
line(28, 152, 1252, 152)
tx(36, 157, 1208, 44, [[('In 2025, Medium and Heavy imports reached ', False, '000000'), (f'{MHU} units ({rp(MHV)})', True, '000000'),
    (f' out of {TU} units ({rp(TV)}). ', False, '000000'), ('Sany Palfinger leads Medium with half of the units; Palfinger leads Heavy with 56%.', True, '000000')]], 14)
line(28, 204, 1252, 204)

# ---------- panel 1: market by class ----------
P1X, P1W = 36, 380
banner(P1X, P1W, 'Market 2025: focus on Medium & Heavy', '595959'); rect(P1X, 240, P1W, 412, None, '595959', 1.5)
tx(P1X + 12, 248, P1W - 24, 34, [[('Market size by class, 2025, ', True, '000000'), ('units vs value', False, '000000')]], 12)
cd = CategoryChartData(); cd.categories = [f'Units ({TU})', f'Value ({rp(TV)})']
for c in CLS: cd.add_series(c, [U[c], V[c] * FX / 1e9])
CX, CY, CW, CH = P1X + 82, 286, 200, 330
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, E(CX), E(CY), E(CW), E(CH), cd).chart
ch.has_legend = False; ch.font.name = FONT; ch.font.size = Pt(9)
pl = ch.plots[0]; pl.gap_width = 55; pl.overlap = 100
ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
ca = ch.category_axis; ca.major_tick_mark = XL_TICK_MARK.NONE; ca.format.line.color.rgb = rgb('000000')
ca.tick_labels.font.size = Pt(10); ca.tick_labels.font.bold = True; ca.tick_labels.font.name = FONT
pa = ch._chartSpace.find('.//' + qn('c:plotArea')); lay = pa.find(qn('c:layout'))
if lay is None:
    lay = etree.Element(qn('c:layout')); pa.insert(0, lay)
lay.append(etree.fromstring('<c:manualLayout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:layoutTarget val="inner"/>'
    '<c:xMode val="edge"/><c:yMode val="edge"/><c:x val="0"/><c:y val="0"/><c:w val="1"/><c:h val="0.9"/></c:manualLayout>'))
bc = ch._chartSpace.find('.//' + qn('c:barChart'))
bc.findall(qn('c:axId'))[0].addprevious(etree.fromstring('<c:serLines xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:spPr><a:ln w="6350"><a:solidFill><a:srgbClr val="7F7F7F"/></a:solidFill></a:ln></c:spPr></c:serLines>'))
for ser, c in zip(ch.series, CLS):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(COL[c]); ser.format.line.color.rgb = rgb('FFFFFF')
    dl = ser.data_labels; dl.font.size = Pt(9); dl.font.name = FONT; dl.show_value = False
    for k, t in enumerate([f'{c} {U[c]}\n{U[c] / TU:.0%}', f'{rp(V[c])}\n{V[c] / TV:.0%}']):
        tf = ser.points[k].data_label.text_frame; tf.text = t
        for p in tf.paragraphs:
            for r in p.runs: r.font.bold = True; r.font.color.rgb = rgb('FFFFFF' if c in ('Medium', 'Heavy') else '595959')
PH = CH * 0.9; fy = lambda f: CY + PH * (1 - f)
ls = (U['Light'] + U['Small']) / TU; lsv = (V['Light'] + V['Small']) / TV
# left bracket: Light+Small excluded
bx = CX - 8
line(bx, fy(ls) + 2, bx, fy(0) - 2, GREY, 0.75); line(bx, fy(ls) + 2, bx + 5, fy(ls) + 2, GREY, 0.75); line(bx, fy(0) - 2, bx + 5, fy(0) - 2, GREY, 0.75)
tx(P1X + 8, (fy(ls) + fy(0)) / 2 - 26, bx - P1X - 14, 52, [[('Light +', False, GREY)], [('Small', False, GREY)], [('excluded', True, GREY)]], 9, PP_ALIGN.RIGHT)
line(bx, fy(1) + 2, bx, fy(ls) - 2, ACC, 1); line(bx, fy(1) + 2, bx + 5, fy(1) + 2, ACC, 1); line(bx, fy(ls) - 2, bx + 5, fy(ls) - 2, ACC, 1)
tx(P1X + 8, (fy(1) + fy(ls)) / 2 - 20, bx - P1X - 14, 40, [[(f'{MHU / TU:.0%}', True, ACC)], [(f'{MHU} units', False, '000000')]], 10, PP_ALIGN.RIGHT)
rx = CX + CW + 6
line(rx, fy(1) + 2, rx, fy(lsv) - 2, ACC, 1); line(rx, fy(1) + 2, rx - 5, fy(1) + 2, ACC, 1); line(rx, fy(lsv) - 2, rx - 5, fy(lsv) - 2, ACC, 1)
tx(rx + 6, (fy(1) + fy(lsv)) / 2 - 20, P1X + P1W - rx - 10, 40, [[(f'{MHV / TV:.0%}', True, ACC)], [(rp(MHV), False, '000000')]], 10)

# ---------- arrow (funnel) ----------
def arrow(x):
    a = s.shapes.add_shape(MSO_SHAPE.CHEVRON, E(x), E(420), E(18), E(36)); a.fill.solid(); a.fill.fore_color.rgb = rgb(ACC); a.line.fill.background()
arrow(421); arrow(830)

# ---------- panels 2 & 3: treemaps ----------
def brand_panel(x, w, cls, brands, tm):
    banner(x, w, f'{cls} class ({tm}): who fills it', ACC); rect(x, 240, w, 412, None, ACC, 1.5)
    tx(x + 12, 248, w - 24, 34, [[(f'{cls} class market, 2025, ', True, '000000'), ('n of units by brand', False, '000000')]], 12)
    tx(x + 12, 270, w - 24, 18, [[(f'{U[cls]} units · {rp(V[cls])} · {V[cls] / TV:.0%} of market value', False, GREY)]], 10)
    treemap(x + 12, 296, w - 24, 230, brands)
    tx(x + 12, 538, w - 24, 18, [[(f'Triatra {cls} market share, ', True, '000000'), ('quantity of sales', False, '000000')]], 11)
    rect(x + 12, 562, 10, 10, 'FF0000' if False else ACC, ACC, 0.5); tx(x + 26, 559, 60, 16, 'Triatra', 9)
    rect(x + 84, 562, 10, 10, 'D9D9D9', 'D9D9D9', 0.5); tx(x + 98, 559, 100, 16, 'Rest of market', 9)
    b = rect(x + 12, 584, w - 24, 30, 'F2F2F2', LG, 0.75); tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = f'Triatra units sold in 2025: to be added  ·  market = {U[cls]} units'
    r.font.size = Pt(10); r.font.italic = True; r.font.name = FONT; r.font.color.rgb = rgb(GREY)
    lead = brands[0]
    tx(x + 12, 620, w - 24, 28, [[(f'{lead[0]} holds {lead[1]} of {U[cls]} units ({lead[1] / U[cls]:.0%})', True, ACC)]], 11)
brand_panel(452, 374, 'Medium', MED, '25–45 tm'); brand_panel(862, 382, 'Heavy', HEV, '> 45 tm')

# ---------- footer ----------
tx(36, 662, 1208, 16, [[('Source: ', True, '000000'), ('Indonesia import records 2025, knuckle boom (folding) cranes; value converted from USD at assumed Rp 16.000/USD (Bio = billion).', False, '000000')]], 10)
tx(36, 680, 1208, 16, [[('Note: ', True, GREY), ('Brand units for top 3 per class from earlier data cut; Others = 2025 class total minus top 3. Light (< 8 tm) and Small (8–25 tm) excluded from focus.', False, GREY)]], 10)
line(28, 700, 1252, 700, ACC, 1)
P.save('Slide_2025_Medium_Heavy_Funnel.pptx'); print(TU, rp(TV), MHU, rp(MHV), MED, HEV)
