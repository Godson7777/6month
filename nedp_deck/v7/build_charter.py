"""Project Charter slide (tangible KPIs and benefits), in the user's charter layout."""
import sys
sys.path.insert(0, '.')
import crane_data as d, crane_map as cm
from cr_common import *
from pptx.util import Pt

# ---------- numbers ----------
U, V = d.by_year(); TU = sum(U.values()); TV = sum(V.values()) / 1e9
avg_u = sum(U[(y, c)] for y in (2023, 2024, 2025) for c in d.CLS) / 3; avg_v = sum(V[(y, c)] for y in (2023, 2024, 2025) for c in d.CLS) / 3 / 1e9
hy = [r for r in d.RAW if r[0] == 'HYVA']; hy_u = sum(r[8] for r in hy); hy_v = sum(r[10] for r in hy) * d.FX / 1e9
hy_avg = sum(r[8] for r in hy if r[4] in (2023, 2024, 2025)) / 3
allr = [(c, t) for c in d.CLS for t in cm.tiers(c)]
win = [(c, t) for c, t in allr if t['ch'] and t['ch']['p'] < t['top']['p']]
wu = sum(t['units'] for _, t in win)
GPR = .35
PLAN = [('SQ3.2ZK2', 5), ('SQZ325.4', 2), ('KSQZ365.4', 2), ('GSQZ200.3', 1), ('FBR450 A4', 5)]   # 10 XCMG + 5 F.lli Ferrari FBR450 A4   # 15 cranes, round plan
chp = {t['ch']['type']: t['ch']['p'] for _, t in win}
rev = sum(chp[m] / (1 - GPR) * n for m, n in PLAN) / 1e9
gp = rev * GPR
gaps = [(t['top']['p'] - t['ch']['p']) / 1e6 for _, t in win]
print(f'rev {rev:.2f} gp {gp:.2f} gaps {min(gaps):.0f}-{max(gaps):.0f} hyva {hy_u} {hy_v:.1f} avg {hy_avg:.0f}')

deck = Deck()
C = deck.slide('Project Charter/', 'What is the background,', 'the objective and the benefit?', 'x', [('x', False)], title_w=600)
s = C.s
for sh in list(s.shapes):
    if sh.name == 'Table 2': sh._element.getparent().remove(sh._element)
FS = 11.5
def section(x, y, w, title, h=24):
    C.box(x, y, w, h, G); C.tx(x + 10, y + 3, w - 20, h - 4, [[(title, True, 'FFFFFF', 13)]], anchor=MSO_ANCHOR.MIDDLE)
    return y + h + 5
def bullets(x, y, w, items, sz=FS, h=None):
    paras = []
    for it in items:
        if isinstance(it, str): it = [(it, False)]
        paras.append([('•  ', False, G, sz)] + [(t, b, '000000', sz) for t, b in it])
    b = C.tx(x + 10, y, w - 20, h or 16 * len(items), paras)
    for p in b.text_frame.paragraphs: p.space_after = Pt(2)
    return b

# ================= left column =================
FS = 10.5
def blist(x, y, w, items, sz=FS, lh=17):
    cpl = (w - 34) / (sz * .6)
    n = sum(1 + int(len(''.join(t for t, _ in it) if not isinstance(it, str) else it) / cpl) for it in items)
    h = n * sz * 1.62 + len(items) * 2.7
    bullets(x, y, w, items, sz, h); return y + h + 5
LX, LW = 36, 610; y = 150
y = section(LX, y, LW, 'Project name')
C.tx(LX + 10, y, LW - 20, 18, [[('Indonesia Folding Crane Market Strategy: one challenger crane for every tm class', False, '000000', FS)]]); y += 22
y = section(LX, y, LW, 'Background')
y = blist(LX, y, LW, [[('Hyva has stopped making cranes: ', True), (f'{hy_u} Hyva cranes (Rp {num(hy_v)} Bio) sold in 2023 – 2026 need a new brand', False)],
                      [(f'Indonesia imports ~{avg_u:.0f} folding cranes a year (Rp {num(avg_v)} Bio) across all four classes', False)]])
y = section(LX, y, LW, 'Objective')
C.tx(LX + 10, y, LW - 20, 22, [[('Sell 15 folding cranes by December 2027', True, DG, 14)]]); y += 24
y = blist(LX, y, LW, [[('4C: ', True), (f'map all {TU} imported folding cranes (Rp {num(TV)} Bio, 2023 – 2026) by class, brand, model', False)],
                      [('STP: ', True), (f'split into {len(allr)} tm classes and find the top 1 model in each', False)],
                      [('4P: ', True), (f'set a cheaper F.lli Ferrari or XCMG challenger per tm class ({len(win)} of {len(allr)} done)', False)]])
y = section(LX, y, LW, 'Project Scope')
y = blist(LX, y, LW, [[('In: ', True), ('folding cranes, Light to Heavy; price and specification; imports 2023 – 2026', False)],
                      [('Out: ', True), ('telescopic, grabber and forestry cranes; Hyva; Zoomlion', False)]])
y = section(LX, y, LW, 'KPI / Metric')
kpi = [('Units sold', '15 cranes', 'by Dec 2027'),
       ('Revenue', f'≈ Rp {num(rev)} Bio', f'GP ≈ Rp {num(gp)} Bio (35%)'),
       ('Cost vs leader', f'Rp {min(gaps):.0f} – {max(gaps):.0f} Jt lower', f'per unit, in {len(win)} of {len(allr)} tm classes'),
       ('Lead time', '≤ 14 weeks', 'stock in Malaysia (TBC)')]
tb = s.shapes.add_table(len(kpi), 3, E(LX), E(y), E(LW), E(19 * len(kpi))).table
for j, w in enumerate([150, 200, 260]): tb.columns[j].width = E(w)
for i, row in enumerate(kpi):
    tb.rows[i].height = E(19)
    for j, v in enumerate(row):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(LG if j == 1 else 'FFFFFF'); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = E(8); c.margin_top = c.margin_bottom = E(0)
        r = c.text_frame.paragraphs[0].add_run(); r.text = v; r.font.name = FONT; r.font.size = Pt(FS if j < 2 else FS - 1)
        r.font.bold = j == 1; r.font.color.rgb = rgb(DG if j == 1 else '000000')
y += 19 * len(kpi) + 8
y = section(LX, y, LW, 'Support Requirements')
y = blist(LX, y, LW, [[('Business Consultant: ', True), ('leads, customer visits  ·  ', False), ('Marcom: ', True), ('sales kit, exhibitions', False)],
                      [('Application Engineering: ', True), ('tm class, specs  ·  ', False), ('UTPE: ', True), ('quotation, mounting, install', False)]])

# ================= right column =================
RX, RW = 666, 576; y = 150
y = section(RX, y, RW, 'Project Timeframe')
months = ['Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec', 'Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
acts = [('Market study (4C)', range(0, 3)), ('STP & challenger choice', range(3, 5)), ('Negotiate F.lli Ferrari & XCMG', range(4, 7)),
        ('Sales kit with Marcom', range(5, 7)), ('Exhibition & customer visits', range(7, 13)), ('Quotation → PO → delivery', range(7, 19)),
        ('15 cranes sold', range(18, 19))]
AW = 168; NM = len(months); CW = (RW - AW) / NM; RHG = 17
def cellbox(x, y, w, h, fillc, text='', bold=False, sz=9, al=PP_ALIGN.CENTER, color='000000'):
    b = C.box(x, y, w, h, fillc, 'D9D9D9', lw=.5); tf = b.text_frame
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, E(3) if m in ('margin_left', 'margin_right') else 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    if text:
        p = tf.paragraphs[0]; p.alignment = al; r = p.add_run(); r.text = text
        r.font.name = FONT; r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = rgb(color)
cellbox(RX, y, AW, 2 * RHG, 'F2F2F2', 'Activity', True, 9.5, PP_ALIGN.LEFT)
cellbox(RX + AW, y, 7 * CW, RHG, 'F2F2F2', '2026', True, 9); cellbox(RX + AW + 7 * CW, y, 12 * CW, RHG, 'F2F2F2', '2027', True, 9)
for j, m in enumerate(months): cellbox(RX + AW + j * CW, y + RHG, CW, RHG, 'F2F2F2', m[0], False, 8)
for i, (a, rg) in enumerate(acts):
    last = i == len(acts) - 1; yy = y + (2 + i) * RHG
    cellbox(RX, yy, AW, RHG, 'FFFFFF', a, last, 9, PP_ALIGN.LEFT, DG if last else '000000')
    for j in range(NM):
        cellbox(RX + AW + j * CW, yy, CW, RHG, (DG if last else '3CC46B') if j in rg else 'FFFFFF', '★' if (last and j in rg) else '', True, 8, PP_ALIGN.CENTER, 'FFFFFF')
y += RHG * (2 + len(acts)) + 8
y = section(RX, y, RW, 'Benefit of Project')
y = blist(RX, y, RW, [[(f'≈ Rp {num(rev)} Bio revenue, Rp {num(gp)} Bio GP', True), (' from 15 cranes by end-2027', False)],
                      [(f'Replaces Hyva: ~{hy_avg:.0f} cranes a year', True), (' get a new brand', False)],
                      [(f'Our cost is Rp {min(gaps):.0f} – {max(gaps):.0f} Jt below the leader', True), (' per unit: room for 35% GP', False)],
                      [('Lead time ≤ 14 weeks', True), (' (China / Malaysia) vs 25 – 27 weeks (Italy / Brazil)', False)]], sz=10.5, lh=16)
y = section(RX, y, RW, 'Project Team')
team = [('Executive Sponsor', 'David'), ('Advisor', 'Ceisar Centiaga'), ('Owner', 'Ricardo Tanada'), ('Coordinator', 'Ignatius Indrawan Sumawinata'),
        ('Project Lead', 'Tiberias Krisgaharu Simu'), ('Team Member', 'Business Consultant, Application Engineering, Marcom, UTPE')]
C.tx(RX + 10, y, 130, 15 * len(team), [[(a, False, '000000', 10)] for a, _ in team])
C.tx(RX + 130, y, RW - 140, 15 * len(team), [[(': ' + b, False, '000000', 10)] for _, b in team]); y += 16 * len(team) + 8
y = section(RX, y, RW, 'Project Approval')
roles = [('Lead', 'Tiberias K. Simu'), ('Coordinator', 'Ignatius Indrawan S.'), ('Owner', 'Ricardo Tanada'), ('Advisor', 'Ceisar Centiaga'), ('Executive Sponsor', 'David')]
bw = (RW - 4 * 6) / 5
for k, (role, name) in enumerate(roles):
    x = RX + k * (bw + 6)
    C.tx(x, y, bw, 13, [[(role, True, '000000', 9)]]); C.tx(x, y + 12, bw, 11, [[('Date:', False, '7F7F7F', 8)]])
    C.box(x, y + 25, bw, 26, 'FFFFFF', 'BFBFBF', lw=.75)
    C.box(x, y + 53, bw, 16, 'FFFFFF', 'BFBFBF', lw=.75); C.tx(x, y + 55, bw, 13, [[(name, False, '000000', 8.5)]], al=PP_ALIGN.CENTER)
print('right bottom', y + 70, 'left ok')
C.tx(36, 700, 1206, 14, [[(f'Revenue plan, 15 cranes: 10 XCMG (SQ3.2ZK2 ×5, SQZ325.4 ×2, KSQZ365.4 ×2, GSQZ200.3 ×1) + 5 F.lli Ferrari FBR450 A4; selling price = cost ÷ (1 − 35% GP). '
                          'Prices and units from the crane market slides (imports 2023 – 2026, F.lli Ferrari price list June 2026). Lead time per F.lli Ferrari price list; Malaysia stock to be confirmed.', False, '7F7F7F', 7)]])
deck.save('v7/Project_Charter.pptx'); print('ok')
