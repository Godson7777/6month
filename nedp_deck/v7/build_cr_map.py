"""Crane deck: competitor map per class (simplified price comparison table) + challenger column. Usage: python3 v7/build_cr_map.py Light"""
import copy, sys, collections, re
sys.path.insert(0, '.')
import crane_data as d, crane_map as cm
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string; FONT = 'Neue Haas Grotesk Text Pro'
num = lambda x, n=1: f'{x:,.{n}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
jt = lambda x: f"Rp {num(x / 1e6, 0)} Jt"
G, DG, LG = '1E9E57', '0B4D2A', 'E1F9EC'
CLS = sys.argv[1]; RANGE = {'Light': '≤ 8 tm', 'Small': '> 8 – 25 tm', 'Medium': '> 25 – 45 tm', 'Heavy': '> 45 tm'}[CLS]
LOGO = {'AMCO VEBA': 'amco_veba', 'FASSI': 'fassi', 'HIAB': 'hiab', 'PALFINGER': 'palfinger', 'SANY PALFINGER': 'sany_palfinger', 'XCMG': 'xcmg', 'FERRARI': 'ferrari'}
T = cm.tiers(CLS)
cu = collections.Counter(); 
for t in T:
    for x in t['models']: cu[x['brand']] += x['q']
BR = [b for b, _ in cu.most_common() if b != 'FERRARI']
TU = sum(t['units'] for t in T); TA = sum(t['amt'] for t in T)
main = max(T, key=lambda t: t['units'])

P = Presentation('v7/6M_V9.pptx'); L = P.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(L)))):
    if i != 13: P.part.drop_rel(sid.rId); L.remove(sid)
src = P.slides[0]; s = P.slides.add_slide(src.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in src.shapes:
    if sh.name in ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2'):
        s.shapes._spTree.append(copy.deepcopy(sh._element))
P.part.drop_rel(L[0].rId); L.remove(L[0])
def set_runs(par, text):
    rs = par.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
SH = {sh.name: sh for sh in s.shapes}
from lxml import etree
grp = SH['Group 31']; xf = grp._element.find('.//' + qn('a:xfrm'))
for tag in ('a:ext', 'a:chExt'): xf.find(qn(tag)).set('cx', str(E(1242 - 36.45)))
for c in grp.shapes:
    if 'Medium' in c.text_frame.text: set_runs(c.text_frame.paragraphs[0], f'{CLS} Class Competitor Map/'); c.width = E(860)
    else:
        ps = c.text_frame.paragraphs; set_runs(ps[0], 'Who leads each tm class?'); r0_ = ps[0].runs[0]._r
        br = etree.SubElement(ps[0]._p, qn('a:br')); ps[0]._p.remove(br); r0_.addnext(br)
        if r0_.find(qn('a:rPr')) is not None: br.append(copy.deepcopy(r0_.find(qn('a:rPr'))))
        r1_ = copy.deepcopy(r0_); br.addnext(r1_); r1_.find(qn('a:t')).text = 'Which crane beats it on price?'
        c.left = E(900); c.width = E(342)
t2 = SH['Table 2'].table
mt, mc = main['top'], main['ch']
head = (f"{CLS}: {mt['nice']} {mt['type']} leads; " + (f"{mc['nice']} {mc['type']} is {jt(mt['p'] - mc['p'])} cheaper per unit" if mc else 'no cheaper crane of the same size yet'))
set_runs(t2.cell(0, 0).text_frame.paragraphs[0], head)
p1 = t2.cell(1, 0).text_frame.paragraphs[0]; r0 = p1.runs[0]
for r in p1.runs[1:]: r._r.getparent().remove(r._r)
parts = [(f'{CLS} class ({RANGE}): ', False), (f'{TU} units worth Rp {num(TA / 1e9)} Bio', True), (' in Jan 2023 – 14 Aug 2026. ', False),
         (f'{main["tier"][0]}–{main["tier"][1]} tm takes {main["units"] / TU:.0%} of units', True), ('. Each row compares the tier leader (by value) with the cheapest crane of about the same size.', False)]
r0.text = parts[0][0]; prev = r0._r
for _ in parts[1:]:
    nr = copy.deepcopy(r0._r); prev.addnext(nr); prev = nr
for r, (t, b) in zip(p1.runs, parts): r.text = t; r.font.bold = b

# ---------- table ----------
cols = ['tm class', 'Our challenger', 'F.lli Ferrari'] + BR
W0, W1, W2 = 92, 170, 150; WB = (1206 - W0 - W1 - W2) / len(BR)
widths = [W0, W1, W2] + [WB] * len(BR)
X0, Y0 = 36, 246
tb_shape = s.shapes.add_table(1 + len(T), len(cols), E(X0), E(Y0), E(sum(widths)), E(300)); tb = tb_shape.table
for j, w in enumerate(widths): tb.columns[j].width = E(w)
tb.rows[0].height = E(48)
RH = max(80, (648 - Y0 - 48) / len(T))
for i in range(1, 1 + len(T)): tb.rows[i].height = E(RH)
def fill(c, color): c.fill.solid(); c.fill.fore_color.rgb = rgb(color)
def write(c, paras, al=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = c.text_frame; c.vertical_anchor = anchor; c.margin_left = c.margin_right = E(6); c.margin_top = E(5); c.margin_bottom = E(3)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.alignment = al
        for t, bold, color, sz in para:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
# header
for j, name in enumerate(cols):
    c = tb.cell(0, j); fill(c, G if j == 1 else DG)
    if j == 0: write(c, [[('tm class', True, 'FFFFFF', 10.5)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    elif j == 1: write(c, [[('Our challenger', True, 'FFFFFF', 11)], [('cheapest of the same size', False, 'E1F9EC', 8.5)]], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    elif j == 2: write(c, [[('', False, 'FFFFFF', 8)], [('', False, 'FFFFFF', 8)], [('price list + imports', False, 'E1F9EC', 8.5)]], PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM)
    else: write(c, [[('', False, 'FFFFFF', 8)], [('', False, 'FFFFFF', 8)], [(f'{cu[name]} units', False, 'E1F9EC', 8.5)]], PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM)
FS = 11 if len(T) <= 2 else 10 if len(T) <= 4 else 9
def model_lines(x, top=False):
    lines = ([[('TOP 1', True, G, FS - 1)]] if top else []) + [
             [(f"{x['type']}", True, '000000', FS + .5)],
             [(f"• {x['q']} unit{'s' if x['q'] > 1 else ''}", False, '000000', FS - .5)],
             [(f"• Rp {num(x['amt'] / 1e9, 1 if x['amt'] >= 1e9 else 2)} Bio total", False, '000000', FS - .5)],
             [(f"• {jt(x['p'])} /unit", False, '000000', FS - .5)], [('', False, '000000', 4)]]
    return lines
for i, t in enumerate(T, start=1):
    lo, hi = t['tier']; lab = f'≤ {hi} tm' if lo == 0 else f'> {lo} – {hi} tm'
    c = tb.cell(i, 0); fill(c, 'F2F2F2'); write(c, [[(lab, True, G, FS + 2)], [(f"{t['units']} units", False, '000000', FS)], [(f"Rp {num(t['amt'] / 1e9)} Bio", False, '000000', FS)]], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    ch = t['ch']; c = tb.cell(i, 1); fill(c, LG if ch else 'F2F2F2')
    if ch:
        srcl = 'F.lli Ferrari price list' if ch.get('src') else f"import data, {ch['q']} unit{'s' if ch['q'] > 1 else ''}"
        write(c, [[(f"{ch['nice']} {ch['type']}", True, DG, FS + 1)],
                  [(f"• {jt(ch['p'])} /unit", False, '000000', FS)], [('', False, '000000', 4)], [(f"{jt(t['top']['p'] - ch['p'])} cheaper", True, G, FS + 2)],
                  [(f"• vs {t['top']['nice']} {t['top']['type']}, {jt(t['top']['p'])} /unit", False, '595959', FS - 1.5)], [('', False, '000000', 4)], [(srcl, False, '7F7F7F', FS - 2.5)]], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    else:
        write(c, [[('No cheaper crane of the same size', False, '595959', 9)]], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    # Ferrari column: imports + price list (cheapest variant per base model)
    c = tb.cell(i, 2); fill(c, 'FFFFFF'); paras = []
    for x in [m for m in t['models'] if m['brand'] == 'FERRARI']:
        paras += model_lines(x, x is t['top'])
    base = {}
    for f in t['fer']:
        k = f['type'].split()[0]
        if k not in base or f['p'] < base[k]['p']: base[k] = f
    for f in sorted(base.values(), key=lambda f: f['tm']):
        paras += [[(f"{f['type']}", True, '000000', FS + .5)], [(f"• {jt(f['p'])} /unit", False, '000000', FS - .5)], [('', False, '000000', 4)]]
    write(c, paras or [[('–', False, '7F7F7F', 9)]])
    for j, b in enumerate(BR, start=3):
        c = tb.cell(i, j); ms = [m for m in t['models'] if m['brand'] == b]
        istop = any(m is t['top'] for m in ms); fill(c, 'FFF7D6' if istop else 'FFFFFF')
        paras = []
        for m in ms: paras += model_lines(m, m is t['top'])
        write(c, paras or [[('', False, '000000', 8)]])
# logos in header
for j, name in enumerate(cols):
    if j < 2: continue
    key = 'FERRARI' if j == 2 else name; x = X0 + sum(widths[:j]); w = widths[j]
    plate = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x + 8), E(Y0 + 5), E(w - 16), E(24)); plate.fill.solid(); plate.fill.fore_color.rgb = rgb('FFFFFF'); plate.line.fill.background(); plate.shadow.inherit = False
    if key in LOGO:
        f = f'assets/logos/brands/{LOGO[key]}.png'; iw, ih = Image.open(f).size; k = min((w - 28) / iw, 18 / ih)
        s.shapes.add_picture(f, E(x + w / 2 - iw * k / 2), E(Y0 + 17 - ih * k / 2), E(iw * k), E(ih * k))
    else:
        tf = plate.text_frame; tf.margin_top = tf.margin_bottom = 0; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = d.NICE.get(key, key.title()); r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = rgb(DG); r.font.name = FONT
# legend + notes
y = 660
for k, (col, lab) in enumerate([(LG, 'Our challenger: cheaper than the tier leader, at most 2 tm smaller'), ('FFF7D6', 'Tier leader (TOP 1 = highest import value in the tier)')]):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(36 + k * 470), E(y), E(12), E(12)); b.fill.solid(); b.fill.fore_color.rgb = rgb(col); b.line.color.rgb = rgb('BFBFBF'); b.shadow.inherit = False
    tb_ = s.shapes.add_textbox(E(54 + k * 470), E(y - 1), E(440), E(16)); tb_.text_frame.margin_left = tb_.text_frame.margin_top = 0; r = tb_.text_frame.paragraphs[0].add_run(); r.text = lab; r.font.size = Pt(9.5); r.font.name = FONT
tn = s.shapes.add_textbox(E(36), E(680), E(1206), E(34)); tf = tn.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run(); r.font.size = Pt(6.5); r.font.name = FONT; r.font.color.rgb = rgb('7F7F7F')
r.text = ('Source: Indonesia import records (Crane Market GVW Pairing v11, Noir), Jan 2023 – 14 Aug 2026; F.lli Ferrari price list June 2026 (Rev1), sheet 25.06. '
          'Competitor price = average import unit price of the model (total value ÷ total units) × Rp 17.803/USD (JISDOR 23 Sep 2026) × (1 + import duty: 0% China/ACFTA, 5% others) × (1 + PPN 11% + PPh 22 2.5%); crane only, before distributor margin. '
          'F.lli Ferrari price = (column K − column I) × (1 − 13.6% TSP GP): crane only, no cargo deck, no install, no warranty; taxes included. '
          'Units and amount = import value before tax. Model variants merged (e.g. SPK8500 = SPK8500, A, B). tm class per Application Engineering (rounded max lifting moment). '
          'Folding cranes only; Hyva, Zoomlion, grabbers, telescopic cranes and invalid-price records excluded.')
P.save(f'v7/CR_Map_{CLS}.pptx'); print('ok', head)
