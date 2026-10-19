"""One slide: Pre Sales (Demand Check, Preliminary Drawing Review, Request Quote to UTE) with photos, in the user's deck style."""
import copy, os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string; FONT = 'Neue Haas Grotesk Text Pro'
P = Presentation('v7/6M_V9.pptx'); L = P.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(L)))):
    if i != 13: P.part.drop_rel(sid.rId); L.remove(sid)
src = P.slides[0]
s = P.slides.add_slide(src.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in src.shapes:
    if sh.name in ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2'):
        s.shapes._spTree.append(copy.deepcopy(sh._element))
P.part.drop_rel(L[0].rId); L.remove(L[0])
def set_runs(par, text):
    rs = par.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
SH = {sh.name: sh for sh in s.shapes}
for c in SH['Group 31'].shapes:
    set_runs(c.text_frame.paragraphs[0], 'Pre Sales/' if 'Medium' in c.text_frame.text else 'What do I do before the PO comes in?')
t2 = SH['Table 2'].table
set_runs(t2.cell(0, 0).text_frame.paragraphs[0], 'Before the PO, I check the demand, review the drawing and get the price from UTE')
set_runs(t2.cell(1, 0).text_frame.paragraphs[0], 'Three steps that turn a customer need into a quote the Business Consultant can send.')
def tx(x, y, w, h, paras, sz, color='000000', bold=False, al=PP_ALIGN.LEFT):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = al; p.space_after = Pt(3)
        r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.name = FONT; r.font.bold = bold; r.font.color.rgb = rgb(color)
    return b
os.makedirs('/tmp/claude-0/ps', exist_ok=True)
def pic(f, x, y, w, h):
    path = 'assets/evidence/' + f
    if f.endswith('.webp'):
        out = '/tmp/claude-0/ps/' + f.replace('.webp', '.png'); Image.open(path).convert('RGB').save(out); path = out
    iw, ih = Image.open(path).size; k = min(w / iw, h / ih); pw, ph = iw * k, ih * k
    fr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    fr.fill.solid(); fr.fill.fore_color.rgb = rgb('F2F2F2'); fr.line.fill.background(); fr.shadow.inherit = False
    p_ = s.shapes.add_picture(path, E(x + (w - pw) / 2), E(y + (h - ph) / 2), E(pw), E(ph))
    p_.line.color.rgb = rgb('D9D9D9'); p_.line.width = Pt(.5)
STEPS = [
 ('Demand Check', ['Spot which demands are rising in PCR',
                   'Why the customer buys: expansion, new business or replacement (with the Business Consultant)',
                   'Project location, and whether the customer wants it fast or cheap',
                   'Expected price and when the PO may come'],
  [('sc1_a.png', 0, 0, 394, 140), ('sc1_b.webp', 0, 146, 194, 152), ('sc1_c.png', 200, 146, 194, 152)]),
 ('Preliminary Drawing Review', ['What the unit does and how it works',
                   'Main parts that drive the price or the lead time',
                   'Details checked with Application Engineering'],
  [('sc2_a.png', 0, 0, 194, 146), ('sc2_b.png', 200, 0, 194, 146), ('sc2_c.webp', 0, 152, 194, 146), ('sc2_d.png', 200, 152, 194, 146)]),
 ('Request Quote to UTE', ['Drawing, target price, lead time, end customer and delivery date sent to UTE',
                   'Standard GP check (13.6%) before any negotiation',
                   'SLA tracked; escalated when it runs late',
                   'Quote result entered into PCR'],
  [('sc3_a.png', 0, 0, 194, 146), ('sc3_b.png', 200, 0, 194, 146), ('sc3_c.png', 0, 152, 394, 146)]),
]
for i, (head, bullets, pics) in enumerate(STEPS):
    x = 36 + i * 406; W = 394
    hb = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, E(x), E(222), E(W), E(30))
    hb.fill.solid(); hb.fill.fore_color.rgb = rgb('1E9E57'); hb.line.fill.background(); hb.shadow.inherit = False
    tf = hb.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER; r = tf.paragraphs[0].add_run()
    r.text = f'{i + 1}. {head}'; r.font.size = Pt(14); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb('FFFFFF')
    tx(x + 4, 262, W - 8, 112, ['• ' + b for b in bullets], 10.5)
    for f, dx, dy, w, h in pics: pic(f, x + dx, 380 + dy, w, h)
tx(36, 678, 1206, 14, ['Photos: my own work files (PCR in CRM, unit drawings, quote request to UTE, price calculation).'], 7, '6B6B6B')
P.save('v7/Slide_PreSales.pptx'); print('ok', len(P.slides))
