"""Add to the user's 6M_V7 deck: slide 15 Heavy focus (copy of slide 14 layout) and slide 16 4P (Medium + Heavy)."""
import copy, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from lxml import etree
SRC, OUT = 'v7/6M_V7_user2.pptx', 'v7/6M_V9.pptx'
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string
FONT = 'Neue Haas Grotesk Text Pro'
src = open('stp_medium_heavy.py').read(); src = src[:src.index("P = Presentation();")]
G = {'__name__': 'rows'}; exec(compile(src, 'stp_medium_heavy.py', 'exec'), G)
RM, CUM = G['build_rows']('Medium'); RH, CUH = G['build_rows']('Heavy')
rp = lambda x: 'Rp ' + f'{x:,.0f}'.replace(',', '.')
jt = lambda x: 'Rp ' + f'{x / 1e6:,.0f}'.replace(',', '.') + ' Jt'
m = {r['t']: r for r in RM}; h = {r['t']: r for r in RH}
FBR450 = 1350e6; FBR450_TM = 45.5
g45, g50 = FBR450 / h[45]['p'] - 1, FBR450 / h[50]['p'] - 1
hshare = h[45]['ms'] + h[50]['ms']; hunits = h[45]['tq'] + h[50]['tq']

P = Presentation(SRC); S14 = P.slides[13]

def set_runs(par, text):
    rs = par.runs
    if not rs: r = par.add_run(); r.text = text; return
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def set_tf(tf, lines):
    """Replace paragraphs, reusing the i-th original paragraph's format (or the last one)."""
    ps = [p._p for p in tf.paragraphs]; tmpl = [copy.deepcopy(p) for p in ps]
    body = ps[0].getparent()
    for p in ps: body.remove(p)
    for i, line in enumerate(lines):
        np_ = copy.deepcopy(tmpl[min(i, len(tmpl) - 1)]); body.append(np_)
    for par, line in zip(tf.paragraphs, lines): set_runs(par, line)
def add_chart(s, rows, focus, x, y, w, hgt, vmax, anchor_el):
    cd = CategoryChartData()
    cd.categories = [f"{r['t']}–{r['t'] + 5} tm\n{r['type']}" for r in rows]
    cd.add_series('% of value', [round(r['ms'], 4) for r in rows], number_format='0%')
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(x), E(y), E(w), E(hgt), cd)
    anchor_el.addprevious(gf._element)
    ch = gf.chart; ch.has_legend = False; ch.has_title = False
    pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = '0%'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(14); dl.font.bold = True; dl.font.name = FONT
    va = ch.value_axis; va.maximum_scale = vmax; va.minimum_scale = 0; va.visible = False; va.has_major_gridlines = False
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(9 if len(rows) > 4 else 10); ca.tick_labels.font.name = FONT; ca.format.line.color.rgb = rgb('7F7F7F')
    for i, r in enumerate(rows):
        pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb('1E9E57' if r['t'] in focus else 'BFBFBF')
    PX, PY, PW, PH = .02, .26, .96, .58
    pa = ch._chartSpace.find('.//' + qn('c:plotArea')); lay = pa.find(qn('c:layout'))
    if lay is None: lay = etree.SubElement(pa, qn('c:layout')); pa.remove(lay); pa.insert(0, lay)
    lay.append(etree.fromstring(f'<c:manualLayout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                                f'<c:x val="{PX}"/><c:y val="{PY}"/><c:w val="{PW}"/><c:h val="{PH}"/></c:manualLayout>'))
    n = len(rows); catW = w * PW / n; barW = catW / 1.6; x0 = x + w * PX
    i0, i1 = [i for i, r in enumerate(rows) if r['t'] in focus][0], [i for i, r in enumerate(rows) if r['t'] in focus][-1]
    bl = x0 + catW * (i0 + .5) - barW / 2 - 4; br = x0 + catW * (i1 + .5) + barW / 2 + 4
    ytop = y + hgt * PY + hgt * PH * (1 - max(rows[i0]['ms'], rows[i1]['ms']) / vmax) - 28
    return bl, br, ytop

# ---------------- slide 15: Heavy focus (copy of slide 14) ----------------
s = P.slides.add_slide(S14.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in S14.shapes:
    if sh.name in ('Chart 45',): continue
    s.shapes._spTree.append(copy.deepcopy(sh._element))
SH = {sh.name: sh for sh in s.shapes}
for c in SH['Group 31'].shapes:
    if 'Medium' in c.text_frame.text: set_runs(c.text_frame.paragraphs[0], 'Heavy Class: Where to Focus/')
t2 = SH['Table 2'].table
set_runs(t2.cell(0, 0).text_frame.paragraphs[0], f"Heavy: focus on 45–55 tm, where two tiers hold {hshare:.0%} of Heavy value")
set_runs(t2.cell(1, 0).text_frame.paragraphs[0], "Take on Amco Veba V950 and Palfinger PK 53002 with one model, FBR450R A4.")
set_runs(SH['TextBox 36'].text_frame.paragraphs[0], f"{hshare:.0%}")
set_tf(SH['TextBox 37'].text_frame, [f"of Heavy value sits in 45–50 tm ({h[45]['ms']:.0%}) and 50–55 tm ({h[50]['ms']:.0%})"])
set_runs(SH['TextBox 39'].text_frame.paragraphs[0], f"{-g50 * 100:.0f}–{-g45:.0%}")
set_tf(SH['TextBox 40'].text_frame, ["FBR450R A4 is cheaper than the top 1 model in both tiers"])
set_runs(SH['TextBox 42'].text_frame.paragraphs[0], "1")
set_tf(SH['TextBox 43'].text_frame, ["model covers both tiers, only 4 tm smaller than PK 53002"])
set_tf(SH['TextBox 44'].text_frame, ["% of Heavy import value by tier, and top 1 model"])
bl, br, ytop = add_chart(s, RH, (45, 50), 36, 336, 470, 334, .75, SH['Left Brace 46']._element)
bh = 20; bw = br - bl; b = SH['Left Brace 46']
b.left, b.top, b.width, b.height = E((bl + br) / 2 - bh / 2), E(ytop - bh / 2 - bw / 2), E(bh), E(bw)
tb47 = SH['TextBox 47']; tb47.left, tb47.top, tb47.width = E(bl - 50), E(ytop - bh - 40), E(bw + 100)
set_tf(tb47.text_frame, [f"{hshare:.0%} of Heavy value", f"{hunits} of {CUH} units"])
T = SH['Table 4'].table
fbr = f"FBR450R A4 ({int(FBR450_TM + .5)} tm)"
cells = [
 [[f"Amco Veba V950 ({int(h[45]['ttm'] + .5)} tm)", rp(h[45]['p'])], [fbr, rp(FBR450), f"{-g45:.0%} cheaper"],
  ["• Double linkage", "• Power boost", "• Electronic control", "• Magic Touch auto fold"],
  ["• Fast boom extension (RS)", "• Smooth Descent Valve: less boom bounce"]],
 [[f"Palfinger PK 53002 ({int(h[50]['ttm'] + .5)} tm)", rp(h[50]['p'])], [fbr, rp(FBR450), f"{-g50:.0%} cheaper"],
  ["• Double linkage", "• Power boost", "• Fast boom extension", "• Boom damping (SDV = AOS)", "• Electronic control", "• Auto fold"],
  ["• No extra feature", "• Wins on price", "• 4 tm smaller (46 vs 50 tm)"]]]
T.cell(2, 2)._tc.replace(T.cell(2, 2)._tc.txBody, copy.deepcopy(T.cell(1, 2)._tc.txBody))
for i, row in enumerate(cells, start=1):
    for j, lines in enumerate(row): set_tf(T.cell(i, j).text_frame, lines)
# keep table body font consistent in the feature columns
for i in (1, 2):
    for j in (2, 3):
        for p in T.cell(i, j).text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(10); r.font.bold = False; r.font.name = FONT
set_tf(SH['TextBox 50'].text_frame, ['Source: Crane Market GVW Pairing v11 (Noir), Jan 2023 – 14 Aug 2026; Hyva excluded; tier by original max lifting moment. Competitor price = highest import unit price × Rp 17.803/USD × (1 + 5% duty) × (1 + PPN 11% + PPh 22 2.5%). '
    'FBR450R A4 = TSP price list Rev1 Jun 2026 less 13.6% TSP margin and 3% warranty; 45.5 tm, shown rounded. Features: Amco Veba, Palfinger and F.lli Ferrari brochures and earlier feature comparison.'])

# ---------------- slide 16: 4P Medium + Heavy ----------------
s = P.slides.add_slide(S14.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in S14.shapes:
    if sh.name in ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2'):
        s.shapes._spTree.append(copy.deepcopy(sh._element))
SH = {sh.name: sh for sh in s.shapes}
for c in SH['Group 31'].shapes:
    if 'Medium' in c.text_frame.text: set_runs(c.text_frame.paragraphs[0], 'Marketing Mix (4P)/')
    else: set_tf(c.text_frame, ['How do we win in each tier?'])
t2 = SH['Table 2'].table
set_runs(t2.cell(0, 0).text_frame.paragraphs[0], "Win with European features at a lower price, and stock close to Indonesia")
set_runs(t2.cell(1, 0).text_frame.paragraphs[0], f"Three F.lli Ferrari models cover four tiers: {m[30]['ms'] + m[35]['ms']:.0%} of Medium value and {hshare:.0%} of Heavy value.")
cols = [120, 270, 270, 270, 270]; X0, Y0 = 36, 236
rows = [
 ('', 'Medium 30–35 tm', 'Medium 35–40 tm', 'Heavy 45–50 tm', 'Heavy 50–55 tm'),
 ('Target\n(top 1 model)', f"HIAB X-CLX 388 (34 tm)\n{m[30]['ms']:.0%} of Medium value", f"Palfinger PK 41002 (38 tm)\n{m[35]['ms']:.0%} of Medium value",
  f"Amco Veba V950 (46 tm)\n{h[45]['ms']:.0%} of Heavy value", f"Palfinger PK 53002 (50 tm)\n{h[50]['ms']:.0%} of Heavy value"),
 ('Product', "FBR350R A4\n• Magic Touch auto fold\n• Dynamic Load Diagram\n• Electronic control (HIAB is manual)",
  "7441C (44 tm)\n• 6 tm more capacity\n• Same key features as PK 41002",
  "FBR450R A4 (46 tm)\n• Fast boom extension\n• Smooth Descent Valve",
  "FBR450R A4 (46 tm)\n• Same key features as PK 53002\n• 4 tm smaller"),
 ('Price', f"{jt(m[30]['fp'])}\n{-m[30]['gap']:.0%} cheaper than {jt(m[30]['p'])}",
  f"{jt(m[35]['fp'])}\n{-m[35]['gap']:.0%} cheaper than {jt(m[35]['p'])}",
  f"{jt(FBR450)}\n{-g45:.0%} cheaper than {jt(h[45]['p'])}", f"{jt(FBR450)}\n{-g50:.0%} cheaper than {jt(h[50]['p'])}"),
 ('Place', "Ask F.lli Ferrari to keep crane units and spare parts in Malaysia, so we can deliver to Indonesia in under 14 weeks (TBC).", None, None, None),
 ('Promotion', "Sell as a package: crane + flatdeck + installation + 2-year warranty. Show the units at exhibitions with a Marcom sales kit.", None, None, None),
]
RHs = [28, 50, 92, 50, 44, 44]
tb = s.shapes.add_table(len(rows), len(cols), E(X0), E(Y0), E(sum(cols)), E(sum(RHs))).table
for j, w in enumerate(cols): tb.columns[j].width = E(w)
for i, hh in enumerate(RHs): tb.rows[i].height = E(hh)
for i, row in enumerate(rows):
    if row[2] is None: tb.cell(i, 1).merge(tb.cell(i, 4))
    for j, v in enumerate(row):
        if v is None: continue
        c = tb.cell(i, j); c.text = v; c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = E(8); c.margin_top = c.margin_bottom = E(4)
        head = i == 0 or j == 0
        c.fill.solid(); c.fill.fore_color.rgb = rgb('0B4D2A' if i == 0 else ('E6E6E6' if j == 0 else ('E1F9EC' if i == 3 else 'FFFFFF')))
        for pi, p in enumerate(c.text_frame.paragraphs):
            if i == 0: p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = FONT; r.font.size = Pt(12 if i == 0 else 11)
                r.font.color.rgb = rgb('FFFFFF' if i == 0 else ('1E9E57' if (i == 3 and pi == 1) else '000000'))
                r.font.bold = head or (pi == 0 and i in (1, 2, 3) and row[2] is not None)
tx = s.shapes.add_textbox(E(36), E(678), E(1206), E(30)); tf = tx.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run(); r.font.size = Pt(6.5); r.font.name = FONT; r.font.color.rgb = rgb('6B6B6B')
r.text = ('Source: Crane Market GVW Pairing v11 (Noir), Jan 2023 – 14 Aug 2026; Hyva excluded. Competitor price = highest import unit price × Rp 17.803/USD, incl. duty and import tax. '
          'F.lli Ferrari price = TSP price list Rev1 Jun 2026 less 13.6% TSP margin and 3% warranty; 7441C estimate. Malaysia stock and lead time to be confirmed with F.lli Ferrari.')
P.save(OUT); print('saved', OUT, len(P.slides))
