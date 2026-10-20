"""Crane deck slide 1: market size 2023 - 14 Aug 2026, units and value per year by class (native charts)."""
import copy, sys
sys.path.insert(0, '.')
import crane_data as d
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string; FONT = 'Neue Haas Grotesk Text Pro'
num = lambda x, n=1: f'{x:,.{n}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
COL = {'Light': 'CDEFD9', 'Small': '74CC95', 'Medium': '1E9E57', 'Heavy': '0B4D2A'}
LBL = {'Light': 'Light (≤ 8 tm)', 'Small': 'Small (> 8 – 25 tm)', 'Medium': 'Medium (> 25 – 45 tm)', 'Heavy': 'Heavy (> 45 tm)'}

U, V = d.by_year(); YRS = [2023, 2024, 2025, 2026]
TU = sum(U.values()); TV = sum(V.values()) / 1e9
AU = sum(U[(y, c)] for y in YRS[:3] for c in d.CLS) / 3; AV = sum(V[(y, c)] for y in YRS[:3] for c in d.CLS) / 3 / 1e9
YU = sum(U[(2026, c)] for c in d.CLS); YV = sum(V[(2026, c)] for c in d.CLS) / 1e9

P = Presentation('v7/6M_V9.pptx'); L = P.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(L)))):
    if i != 13: P.part.drop_rel(sid.rId); L.remove(sid)
src = P.slides[0]; s = P.slides.add_slide(src.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in src.shapes:
    if sh.name in ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2'):
        s.shapes._spTree.append(copy.deepcopy(sh._element))
P.part.drop_rel(L[0].rId); L.remove(L[0])
def set_runs(par, text):
    rs = par.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
SH = {sh.name: sh for sh in s.shapes}
for c in SH['Group 31'].shapes:
    set_runs(c.text_frame.paragraphs[0], 'Crane Market 2023–2026/' if 'Medium' in c.text_frame.text else 'How big is the market, and how much sells each year?')
t2 = SH['Table 2'].table
set_runs(t2.cell(0, 0).text_frame.paragraphs[0], f'Indonesia imports about {AU:.0f} folding cranes a year, worth Rp {AV:.0f} Bio')
set_runs(t2.cell(1, 0).text_frame.paragraphs[0], f'From Jan 2023 to 14 Aug 2026 the market took {TU} units worth Rp {num(TV)} Bio, and every class sells every year.')
def tx(x, y, w, h, runs, sz, al=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.alignment = al
        for t, bold, c in para:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz if not isinstance(bold, tuple) else bold[1]); r.font.bold = bold if not isinstance(bold, tuple) else bold[0]
            r.font.name = FONT; r.font.color.rgb = rgb(c)
    return b
# KPI strip
KP = [(f'{TU} units', f'Rp {num(TV)} Bio', 'Total imports, Jan 2023 – 14 Aug 2026'),
      (f'{AU:.0f} units', f'Rp {num(AV)} Bio', 'Average per year, 2023 – 2025'),
      (f'{YU} units', f'Rp {num(YV)} Bio', '2026 so far, Jan – 14 Aug')]
for i, (a, b_, c) in enumerate(KP):
    x = 36 + i * 406
    bx = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(234), E(394), E(62))
    bx.fill.solid(); bx.fill.fore_color.rgb = rgb('E1F9EC' if i < 2 else 'F2F2F2'); bx.line.fill.background(); bx.shadow.inherit = False
    tx(x + 14, 240, 366, 30, [[(a, (True, 20), '1E9E57' if i < 2 else '595959'), ('   ·   ', (False, 16), '7F7F7F'), (b_, (True, 20), '1E9E57' if i < 2 else '595959')]], 20)
    tx(x + 14, 272, 366, 18, [[(c, False, '000000')]], 11)
def stacked(x, y, w, h, title, vals, fmt, vmax, tot_fmt, min_lbl):
    tx(x, y, w, 18, [[(title, True, '000000')]], 12, PP_ALIGN.CENTER)
    cd = CategoryChartData(); cd.categories = ['2023', '2024', '2025', '2026*']
    for c in d.CLS: cd.add_series(LBL[c], [vals[(yy, c)] for yy in YRS], number_format=fmt)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, E(x), E(y + 22), E(w), E(h - 22), cd); ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    pl = ch.plots[0]; pl.gap_width = 55; pl.overlap = 100
    bc = ch._chartSpace.find('.//' + qn('c:barChart'))
    sl = etree.SubElement(bc, qn('c:serLines')); bc.remove(sl); bc.find(qn('c:overlap')).addnext(sl)
    sp = etree.SubElement(sl, qn('c:spPr')); ln = etree.SubElement(sp, qn('a:ln'), w='6350'); sf = etree.SubElement(ln, qn('a:solidFill'))
    etree.SubElement(sf, qn('a:srgbClr'), val='7F7F7F')
    va = ch.value_axis; va.visible = False; va.has_major_gridlines = False; va.minimum_scale = 0; va.maximum_scale = vmax
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(11); ca.tick_labels.font.name = FONT; ca.format.line.color.rgb = rgb('7F7F7F')
    for k, c in enumerate(d.CLS):
        se = pl.series[k]; se.format.fill.solid(); se.format.fill.fore_color.rgb = rgb(COL[c])
        se.data_labels.show_value = True; se.data_labels.number_format = fmt; se.data_labels.number_format_is_linked = False
        se.data_labels.position = XL_LABEL_POSITION.CENTER
        se.data_labels.font.size = Pt(11); se.data_labels.font.bold = True; se.data_labels.font.name = FONT
        se.data_labels.font.color.rgb = rgb('FFFFFF' if c in ('Medium', 'Heavy') else '000000')
        for j, yy in enumerate(YRS):
            if vals[(yy, c)] < min_lbl:
                dls = se._element.find(qn('c:dLbls'))
                dl = etree.fromstring(f'<c:dLbl xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:idx val="{j}"/><c:delete val="1"/></c:dLbl>')
                dls.insert(0, dl)
    PX, PY, PW, PH = .04, .04, .92, .86
    pa = ch._chartSpace.find('.//' + qn('c:plotArea')); lay = pa.find(qn('c:layout'))
    if lay is None: lay = etree.SubElement(pa, qn('c:layout')); pa.remove(lay); pa.insert(0, lay)
    lay.append(etree.fromstring(f'<c:manualLayout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                                f'<c:x val="{PX}"/><c:y val="{PY}"/><c:w val="{PW}"/><c:h val="{PH}"/></c:manualLayout>'))
    cy, chh = y + 22, h - 22; catW = w * PW / 4
    for j, yy in enumerate(YRS):
        t = sum(vals[(yy, c)] for c in d.CLS); cx = x + w * PX + catW * (j + .5)
        top = cy + chh * PY + chh * PH * (1 - t / vmax)
        tx(cx - 60, top - 20, 120, 16, [[(tot_fmt(t), True, '000000')]], 12, PP_ALIGN.CENTER)
VB = {k: v / 1e9 for k, v in V.items()}
stacked(36, 312, 590, 330, 'Units imported per year, by class', U, '0', 230, lambda t: f'{t:.0f} units', 4)
stacked(652, 312, 590, 330, 'Import value per year, by class (Rp Bio)', VB, '0', 125, lambda t: f'Rp {num(t)} Bio', 3)
# legend
for i, c in enumerate(d.CLS):
    x = 300 + i * 190
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(651), E(12), E(12)); b.fill.solid(); b.fill.fore_color.rgb = rgb(COL[c]); b.line.fill.background(); b.shadow.inherit = False
    tx(x + 18, 649, 170, 16, [[(LBL[c], False, '000000')]], 11)
tx(36, 674, 1206, 34, [[('Source: Crane Market GVW Pairing v11 (Noir), Indonesia import records Jan 2023 – 14 Aug 2026; folding (knuckle boom) cranes only. '
    'Excluded: all rows in the Noir "Excluded Rows" sheet (telescopic/stiff boom, grabber, forestry, parts, generic models), Hyva, Zoomlion, HIAB 19000XG and 6000XG (grabber), '
    'Tadano TM-Z series (telescopic), and two Palfinger records with invalid prices (USD 2.000 and 2.680 per unit). Value = import value before tax × Rp 17.803/USD (JISDOR 23 Sep 2026). '
    'Class by max lifting moment, rounded, per Application Engineering. Average per year uses the full years 2023 – 2025. *2026: Jan – 14 Aug only.', False, '6B6B6B')]], 6.5)
P.save('v7/CR1_Market_Size.pptx'); print('ok', TU, round(TV, 1), round(AU), round(AV, 1), YU, round(YV, 1))
