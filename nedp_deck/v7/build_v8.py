"""Continue the user's 6M_V7 deck (layout kept): fix numbers/text to the Noir database at Rp 17.803/USD,
then add STP Medium and STP Heavy slides using slide 13's frame."""
import copy, io, re, collections, openpyxl
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
SRC, OUT = 'v7/6M_V7_user.pptx', 'v7/6M_V8.pptx'
DB = 'data/Crane_Market_GVW_Pairing_v11_Noir.xlsx'; FX = 17803
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string
R = list(openpyxl.load_workbook(DB, data_only=True, read_only=True)['8. Pivot Source'].iter_rows(min_row=2, values_only=True))
CN = {'SANY PALFINGER', 'XCMG', 'BOB-LIFT', 'ZOOMLION'}
CLS = ['Light', 'Small', 'Medium', 'Heavy']
def agg(f): return sum(r[8] for r in R if f(r)), sum(r[10] for r in R if f(r))
bio = lambda usd: usd * FX / 1e9
fb = lambda x: f'{x:,.1f}'.replace('.', ',')               # 231,7
U = {c: agg(lambda r, c=c: r[2] == c) for c in CLS}
U25 = {c: agg(lambda r, c=c: r[2] == c and r[4] == 2025) for c in CLS}
TQ, TV = sum(u for u, v in U.values()), sum(v for u, v in U.values())
MQ, MV = U['Medium'][0] + U['Heavy'][0], U['Medium'][1] + U['Heavy'][1]
T25Q, T25V = sum(u for u, v in U25.values()), sum(v for u, v in U25.values())
M25Q, M25V = U25['Medium'][0] + U25['Heavy'][0], U25['Medium'][1] + U25['Heavy'][1]
P = Presentation(SRC)

def rep_tf(tf, old, new):
    """replace text inside a text frame; keeps run formatting (collapses into first run only if text spans runs)."""
    hit = False
    for p in tf.paragraphs:
        for r in p.runs:
            if old in r.text: r.text = r.text.replace(old, new); hit = True
    if not hit:
        for p in tf.paragraphs:
            full = ''.join(r.text for r in p.runs)
            if old in full and p.runs:
                p.runs[0].text = full.replace(old, new)
                for r in p.runs[1:]: r.text = ''
                hit = True
    return hit
def rep_shape(sh, pairs):
    n = 0
    tfs = []
    if sh.has_text_frame: tfs.append(sh.text_frame)
    if getattr(sh, 'has_table', False) and sh.has_table: tfs += [c.text_frame for row in sh.table.rows for c in row.cells]
    if sh.shape_type == 6: 
        for c in sh.shapes: n += rep_shape(c, pairs)
    for tf in tfs:
        for o, nw in pairs: n += rep_tf(tf, o, nw)
    return n
def rep_slide(s, pairs):
    n = sum(rep_shape(sh, pairs) for sh in s.shapes); return n

def set_cache(chart, values_by_series):
    """overwrite cached numbers of each series (keeps formulas/labels), values_by_series: list of lists"""
    for ser, vals in zip(chart.series, values_by_series):
        cache = ser._element.find('.//' + qn('c:val')).find('.//' + qn('c:numCache'))
        for pt in cache.findall(qn('c:pt')):
            pt.find(qn('c:v')).text = repr(float(vals[int(pt.get('idx'))]))
def set_wb_cell(chart, sheet, cell, value):
    wbp = chart.part.chart_workbook.xlsx_part
    wb = openpyxl.load_workbook(io.BytesIO(wbp.blob)); wb[sheet][cell] = value
    bio_ = io.BytesIO(); wb.save(bio_); chart.part.chart_workbook.update_from_xlsx_blob(bio_.getvalue())

# ================= slide 10: 4C / STP / 4P summary numbers =================
mv, hv = U['Medium'][1], U['Heavy'][1]
china_m = agg(lambda r: r[2] == 'Medium' and r[0] in CN)[1] / mv
eu_h = agg(lambda r: r[2] == 'Heavy' and r[0] not in CN)[1] / hv
sp_m = agg(lambda r: r[2] == 'Medium' and r[0] == 'SANY PALFINGER')[1] / mv
pa_h = agg(lambda r: r[2] == 'Heavy' and r[0] == 'PALFINGER')[1] / hv
hy = [agg(lambda r, y=y: r[2] == 'Heavy' and r[4] == y)[0] for y in (2023, 2024, 2025)]
mh = [r for r in R if r[2] in ('Medium', 'Heavy')]
gvw = sum(r[8] for r in mh if (r[14] or 0) >= 24) / sum(r[8] for r in mh)
pk = agg(lambda r: r[2] == 'Heavy' and str(r[1]).startswith('PALFINGER PK 53002'))[0]
n10 = rep_slide(P.slides[9], [
    ('Medium + Heavy = 41% of units, 64% of value', f'Medium + Heavy = {MQ / TQ:.0%} of units, {MV / TV:.0%} of value'),
    ('Heavy units 12 → 17 → 35', f'Heavy units {hy[0]} → {hy[1]} → {hy[2]}'),
    ('Sany Palfinger, 38% of value', f'Sany Palfinger, {sp_m:.0%} of value'),
    ('Palfinger, 66% of value; PK 53002 alone sold 39 of 65 units', f'Palfinger, {pa_h:.0%} of value; PK 53002 alone sold {pk} of {U["Heavy"][0]} units'),
    ('China brands take 73% of Medium value', f'China brands take {china_m:.0%} of Medium value'),
    ('European brands take 85% of Heavy value', f'European brands take {eu_h:.0%} of Heavy value'),
    ('88% go on trucks', f'{gvw:.0%} go on trucks'),
    ('Medium 25 – 45 tm: 196 units, Rp 156 M. Heavy > 45 tm: 65 units, Rp 115 M.',
     f'Medium 25 – 45 tm: {U["Medium"][0]} units, Rp {fb(bio(mv))} Bio. Heavy > 45 tm: {U["Heavy"][0]} units, Rp {fb(bio(hv))} Bio.'),
    ('Medium: Sany Palfinger SPK32080, SPK36080, SPK42502. Heavy: Palfinger PK 53002, PK 76002, PK 88002.',
     'Top 1 model by value. Medium: Amco Veba V825, HIAB X-CLX 388, Palfinger PK 41002, Sany Palfinger SPK42502. Heavy: Amco Veba V950, Palfinger PK 53002.'),
    ('Two classes, ten 5-tm tiers', 'Two classes, eleven 5-tm tiers with imports'),
    ('cheaper where it wins (−23% at 25 – 30 tm, −32% at 45 – 50 tm)', 'cheaper where it wins (−23% at 25 – 30 tm, −25% at 30 – 35 tm, −14% at 35 – 40 tm, −32% at 45 – 50 tm)'),
    ('XCMG: same capacity, lower price (SQZ365.4 −9% vs SPK36080).', 'XCMG: bigger crane at a lower price against Sany Palfinger (GSQZ460.4 −16% vs SPK42502 at 40 – 45 tm).'),
    ('XCMG: Triatra special price on SQZ325.4 and SQZ365.4.', 'XCMG: Triatra special price on GSQZ460.4.'),
])
print('slide 10 replacements', n10)

# ================= slide 12: 2023-2026 =================
s12 = P.slides[11]
L_S = U['Light'][0] + U['Small'][0]; LSV = U['Light'][1] + U['Small'][1]
n12 = rep_slide(s12, [
    ('Indonesian Import Folding Crane 23-26/', 'Indonesian Import Folding Crane 23-26'),
    ('62% (Rp 208,3 Bio out of Rp 337,2 Bio)', f'{MV / TV:.0%} (Rp {fb(bio(MV))} Bio out of Rp {fb(bio(TV))} Bio)'),
    ('39% of total import comes from light and small class', f'{MQ / TQ:.0%} of total import comes from Medium and Heavy class'),
    ('61% of total import comes from light and small class', f'{L_S / TQ:.0%} of total import comes from Light and Small class'),
    ('62% of total value or 208,3 Bio comes from Medium and Heavy class', f'{MV / TV:.0%} of total value or Rp {fb(bio(MV))} Bio comes from Medium and Heavy class'),
    ('38% of total value iscomes from light and small', f'{LSV / TV:.0%} of total value comes from Light and Small class'),
    ('118 Units in Avg', f'{L_S * 12 / 44:.0f} units per year'),
    ('80 Units in Avg', f'{MQ * 12 / 44:.0f} units per year'),
    ('Selling Less but gain more value', 'Medium and Heavy sell fewer units but bring more value'),
])
for sh in s12.shapes:   # second chart title was a copy of the first
    if sh.name == 'TextBox 91': rep_tf(sh.text_frame, 'Imported folding cranes by class per year, 2023 – 2026', 'Share of units vs value by class, 2023 – 2026')
    if sh.name == 'Chart 106':
        ch = sh.chart; newv = [round(bio(U[c][1]), 1) for c in CLS]
        cd = CategoryChartData(); cd.categories = list(ch.plots[0].categories)
        for se, c, v in zip(ch.series, CLS, newv): cd.add_series(se.name, [U[c][0], v, v])
        ch.replace_data(cd)
        for se, c, v in zip(ch.series, CLS, newv):
            tf = se.points[1].data_label.text_frame
            rep_tf(tf, re.search(r'Rp [\d,]+ Bio', tf.text).group(0), f'Rp {fb(v)} Bio')
print('slide 12 replacements', n12)

# ================= slide 13: 2025 =================
s13 = P.slides[12]
n13 = rep_slide(s13, [
    ('99 units (Rp 75,4 Bio)', f'{M25Q} units (Rp {fb(bio(M25V))} Bio)'),
    ('out of 227 units (Rp 112,6 Bio)', f'out of {T25Q} units (Rp {fb(bio(T25V))} Bio)'),
    ('44% of total import comes from light and small class', f'{M25Q / T25Q:.0%} of total import comes from Medium and Heavy class'),
    ('67% of total Vaue comes from Medium and Heavy Class', f'{M25V / T25V:.0%} of total value comes from Medium and Heavy class'),
    ('3 brands holds', '3 brands hold'), ('of Unit', 'of units'), ('18% of units', 'Others 18%'),
    ('XCMG holds 86% of value  and 85% of units', 'XCMG hold 86% of value and 85% of units'),
    ('Palfinger holds 82% of units', 'Palfinger hold 82% of units'),
])
for sh in s13.shapes:
    if sh.name == 'TextBox 74':
        rep_tf(sh.text_frame, '99 Units', f'Rp {fb(bio(M25V))} Bio')
    if sh.name == 'Chart 45':
        set_cache(sh.chart, [[U25[c][0], round(bio(U25[c][1]), 1)] for c in CLS]); set_wb_cell(sh.chart, 'Sheet1', 'B9', FX)
    if sh.name == 'Chart 129':
        def bv(cls, b): return bio(agg(lambda r: r[2] == cls and r[4] == 2025 and r[0] == b)[1])
        top = {'Medium': ['SANY PALFINGER', 'PALFINGER', 'XCMG'], 'Heavy': ['PALFINGER', 'AMCO VEBA', 'XCMG']}
        cols = {}
        for cls in ('Medium', 'Heavy'):
            vals = {b: (bv(cls, b) if b in top[cls] else 0) for b in ['SANY PALFINGER', 'PALFINGER', 'AMCO VEBA', 'XCMG']}
            vals['Others'] = bio(U25[cls][1]) - sum(vals.values()); cols[cls] = vals
        order = ['SANY PALFINGER', 'PALFINGER', 'AMCO VEBA', 'XCMG', 'Others']
        set_cache(sh.chart, [[round(cols['Medium'][b], 1), round(cols['Heavy'][b], 1)] for b in order])
        set_wb_cell(sh.chart, 'Sheet1', 'K2', FX)
print('slide 13 replacements', n13)

# ================= slides 14-15: STP Medium / Heavy (slide 13 frame) =================
src = open('stp_medium_heavy.py').read(); src = src[:src.index("P = Presentation();")]
G = {'__name__': 'rows'}; exec(compile(src, 'stp_medium_heavy.py', 'exec'), G)
build_rows = G['build_rows']; rp = lambda x: 'Rp ' + f'{x:,.0f}'.replace(',', '.')
FONT = 'Neue Haas Grotesk Text Pro'
FRAME = ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2')
def frame_slide(title, question, bold, line2):
    s = P.slides.add_slide(s13.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for sh in s13.shapes:
        if sh.name in FRAME: s.shapes._spTree.append(copy.deepcopy(sh._element))
    for sh in s.shapes:
        if sh.name == 'Group 31':
            for c in sh.shapes:
                tf = c.text_frame; old = tf.text
                rep_tf(tf, old, title if c.name == 'TextBox 32' else question)
        if sh.name == 'Table 2':
            c0, c1 = sh.table.cell(0, 0).text_frame, sh.table.cell(1, 0).text_frame
            r0 = c0.paragraphs[0].runs; r0[0].text = bold; [setattr(r, 'text', '') for r in r0[1:]]
            r1 = c1.paragraphs[0].runs; r1[0].text = line2; [setattr(r, 'text', '') for r in r1[1:]]
            for r in r1: r.font.bold = False
    return s
def tx(s, x, y, w, h, runs, sz, al=PP_ALIGN.LEFT):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = al
    for t, bold, c in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(c)
    return b
XCMG_TIERS = {('Medium', 40), ('Heavy', 50), ('Heavy', 70)}
def stp_content(s, cls, rows, seg, tgt, pos):
    for i, (head, body, fill) in enumerate([('SEGMENTATION', seg, 'F2F2F2'), ('TARGETING', tgt, 'F2F2F2'), ('POSITIONING', pos, 'E1F9EC')]):
        x = 36 + i * 406
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(236), E(394), E(92))
        b.fill.solid(); b.fill.fore_color.rgb = rgb(fill); b.line.fill.background(); b.shadow.inherit = False
        tx(s, x + 12, 244, 370, 14, [(head, True, '1E9E57')], 10)
        tx(s, x + 12, 258, 370, 68, [(body, False, '000000')], 10.5)
    cols = [72, 50, 72, 136, 116, 50, 62, 146, 112, 150, 240]
    n = len(rows); RH = min(60, (676 - 340 - 72) // n)
    tb = s.shapes.add_table(3 + n, len(cols), E(36), E(340), E(sum(cols)), E(72 + RH * n)).table
    for j, w in enumerate(cols): tb.columns[j].width = E(w)
    tb.rows[0].height = E(22); tb.rows[1].height = E(24); tb.rows[2].height = E(26)
    for i in range(3, 3 + n): tb.rows[i].height = E(RH)
    def cell(i, j, text, bold=False, fill='FFFFFF', color='000000', sz=11):
        c = tb.cell(i, j); c.text = text; c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = E(3); c.margin_top = c.margin_bottom = E(1)
        c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs: r.font.name = FONT; r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = rgb(color)
    tb.cell(0, 0).merge(tb.cell(0, len(cols) - 1)); cell(0, 0, f'{cls} Class Crane', True, '0B4D2A', 'FFFFFF', 12)
    tb.cell(1, 0).merge(tb.cell(2, 0)); cell(1, 0, 'TM Class', True, 'E6E6E6', sz=10)
    for a, b_, t in [(1, 2, 'Market Size'), (3, 7, 'Top 1 Competitor (by value)'), (8, 10, 'Challenger')]:
        tb.cell(1, a).merge(tb.cell(1, b_)); cell(1, a, t, True, 'E6E6E6', sz=10)
    for j, t in enumerate(['', 'Qty', f'% of {cls}', 'Brand', 'Type', 'Qty', 'Value share', 'Price/unit', 'Brand', 'Type', 'Price/unit']):
        if j: cell(2, j, t, True, 'E6E6E6', sz=10)
    fill = lambda g: 'EDEDED' if g is None else ('E1F9EC' if g < 0 else ('FFF7C2' if g <= .30 else 'FEEBEC'))
    for i, r in enumerate(rows, start=3):
        for j, v in enumerate([f"{r['t']} – {r['t'] + 5}", str(r['tq']), f"{r['ms']:.0%}", r['brand'], f"{r['type']} ({int(r['ttm'] + .5)} tm)", str(r['q']), f"{r['sh']:.0%}", rp(r['p'])]):
            cell(i, j, v, j in (0, 3), sz=11 if j not in (4, 7) else 10)
        if (cls, r['t']) in XCMG_TIERS:                            # tiers fought with XCMG
            if r['xn']: cb, ct, cp, g = 'XCMG', f"{r['xn']} ({int(r['xtm'] + .5)} tm)", r['xp'], r['xgap']
            else: cb, ct, cp, g = 'XCMG', 'Model TBC', None, None
        else:
            cb, ct, cp, g = ('F.lli Ferrari' if r['fn'] else '—'), (r['fn'] or '—'), r['fp'], r['gap']
        pt = 'No model' if ct == '—' else ('To be quoted by XCMG' if ct == 'Model TBC' else 'Price not yet known') if cp is None else ( f"{rp(cp)}  ({g:+.0%})".replace('-', '−'))
        if cls == 'Heavy' and r['t'] == 60: pt = 'Price not yet known · no XCMG model in 60–65 tm import data'
        cell(i, 8, cb, True, fill(g), sz=10); cell(i, 9, ct, False, fill(g), sz=10); cell(i, 10, pt, False, fill(g), sz=10 if len(pt) < 40 else 9)
    foot = ('Source: Crane Market GVW Pairing v11 (Noir), Jan 2023 – 14 Aug 2026. Tier by original max lifting moment; model variants merged (SPK32080 = SPK32080, B, C). '
            f'% of {cls} = tier import value / {cls} import value. Top 1 = model with the highest import value in the tier; value share = model value / tier value. '
            'Competitor price = highest import unit price × Rp 17.803/USD (JISDOR 23 Sep 2026) × (1 + duty: 0% China ACFTA, 5% Europe) × (1 + PPN 11% + PPh 22 2.5%), before distributor margin. '
            'F.lli Ferrari = TSP price list Rev1 Jun 2026 less 13.6% TSP margin and 3% warranty; *7441C estimate; 990R less Rp 168 Jt cargo deck and install; 9601CR A8 = UTPE quote 15 Sep 2026. '
            'XCMG GSQZ460.4 priced like competitors (highest import unit price, 0% duty); no XCMG 50–75 tm model in import data, so model and price TBC. Hyva excluded (cranes discontinued). tm shown rounded. '
            'Fill: green = challenger cheaper, yellow = up to +30%, red = over +30%, grey = no price or no model.')
    tx(s, 36, 678, 1206, 38, [(foot, False, '6B6B6B')], 6.5)

RM, CUM = build_rows('Medium'); RH_, CUH = build_rows('Heavy')
winM = [r for r in RM if r['gap'] is not None and r['gap'] < 0]; mid = [r for r in RM if r['t'] in (30, 35)]; h = {r['t']: r for r in RH_}
x40 = RM[-1]
s14 = frame_slide('Medium Class: Who to Target', 'Which model holds the money? Who can beat it on price?',
                  'Medium: F.lli Ferrari wins at 25–40 tm; XCMG takes on Sany Palfinger at 40–45 tm',
                  f"F.lli Ferrari is {min(-r['gap'] for r in winM):.0%}–{max(-r['gap'] for r in winM):.0%} cheaper than the value leader at 25–40 tm ({sum(r['ms'] for r in winM):.0%} of Medium value); "
                  f"at 40–45 tm XCMG {x40['xn']} is {-x40['xgap']:.0%} cheaper than SPK42502.")
stp_content(s14, 'Medium', RM,
    f"Medium class split into {len(RM)} tm tiers. 30–40 tm holds {sum(r['ms'] for r in mid):.0%} of Medium value ({sum(r['tq'] for r in mid)} of {CUM} units).",
    'Top 1 model by value in each tier: ' + '; '.join(f"{r['brand']} {r['type']} ({r['t']}–{r['t'] + 5})" for r in RM) + '.',
    f"F.lli Ferrari: European features, {min(-r['gap'] for r in winM):.0%}–{max(-r['gap'] for r in winM):.0%} cheaper than the European leaders at 25–40 tm. "
    f"XCMG {x40['xn']} ({int(x40['xtm'] + .5)} tm): bigger crane, {-x40['xgap']:.0%} cheaper than Sany Palfinger SPK42502 ({int(x40['ttm'] + .5)} tm).")
s15 = frame_slide('Heavy Class: Who to Target', 'Which model holds the money? Who can beat it on price?',
                  'Heavy: win 45–50 tm with F.lli Ferrari; take on Palfinger PK 53002 at 50–55 tm with XCMG',
                  f"50–55 tm is {h[50]['ms']:.0%} of Heavy value and PK 53002 takes {h[50]['sh']:.0%} of it. XCMG has no 50–75 tm crane in the import data yet, so a model and price must be quoted.")
stp_content(s15, 'Heavy', RH_,
    f"Heavy class split into {len(RH_)} tm tiers. 45–55 tm holds {h[45]['ms'] + h[50]['ms']:.0%} of Heavy value ({h[45]['tq'] + h[50]['tq']} of {CUH} units); tiers above 55 tm have 1–4 units each.",
    f"Top 1 model by value: Amco Veba V950 at 45–50 tm ({h[45]['sh']:.0%} of tier value) and Palfinger PK 53002 at 50–55 tm ({h[50]['sh']:.0%}).",
    f"F.lli Ferrari FBR450R A4 is {-h[45]['gap']:.0%} cheaper than V950. XCMG for PK 53002 (50–55 tm) and PK 76002 (70–75 tm): ask XCMG for a model and Triatra price.")
P.save(OUT); print('saved', OUT, len(P.slides._sldIdLst), 'slides')
