"""One-slide org chart: Marcom reports to Ricardo, placed at staff level (Agi's row)."""
import copy
from pptx import Presentation
from pptx.util import Emu
P = Presentation('v7/6M_V9.pptx'); L = P.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(L)))):
    if i != 5: P.part.drop_rel(sid.rId); L.remove(sid)
s = P.slides[0]; px = lambda v: v / 9525; E = lambda v: Emu(round(v * 9525))
S = {sh.name: sh for sh in s.shapes}
MARCOM = {'Connector 17', 'Rounded Rectangle 18', 'TextBox 39', 'TextBox 41'}
K = (1000 - 36) / (1204 - 36); nx = lambda x: 36 + (x - 36) * K
for sh in s.shapes:
    if sh.name in MARCOM or px(sh.top) < 170: continue
    if sh.shape_type == 9:   # line
        bx, by, ex, ey = px(sh.begin_x), px(sh.begin_y), px(sh.end_x), px(sh.end_y)
        sh.begin_x, sh.end_x = E(nx(bx)), E(nx(ex))
    else:
        x, w = px(sh.left), px(sh.width); sh.left, sh.width = E(nx(x)), E(w * K)
# lift Director/Ricardo to make room for the Marcom branch on the stem
for n, dy in [('Rounded Rectangle 9', -6), ('TextBox 2', -6), ('TextBox 4', -6), ('Rounded Rectangle 13', -8), ('TextBox 8', -8), ('TextBox 9', -8)]:
    S[n].top = E(px(S[n].top) + dy)
c12 = S['Connector 12']; c12.begin_y, c12.end_y = E(214), E(226)
stem = S['Connector 16']; sx = px(stem.begin_x); stem.begin_y, stem.end_y = E(266), E(290)
# Marcom box at Agi's row, right of the last column
agi = S['Rounded Rectangle 88']; ay, ah = px(agi.top), px(agi.height)
bx, bw = 1028, 214
box, t1, t2 = S['Rounded Rectangle 18'], S['TextBox 39'], S['TextBox 41']
box.left, box.top, box.width, box.height = E(bx), E(ay - 8), E(bw), E(ah + 16)
t1.left, t1.top, t1.width = E(bx + 6), E(ay - 4), E(bw - 12)
t2.left, t2.top, t2.width, t2.height = E(bx + 6), E(ay + 12), E(bw - 12), E(40)
import copy as _c
from pptx.enum.text import PP_ALIGN
tf = t2.text_frame; p0 = tf.paragraphs[0]; r0 = p0.runs[0]
lines = ['Karina Audina E.', 'Adinda Siti K. (Marcomm)', 'Rahma (Graphic Designer)']
for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
for r in p0.runs[1:]: r._r.getparent().remove(r._r)
r0.text = lines[0]; p0.alignment = PP_ALIGN.CENTER
for ln in lines[1:]:
    np_ = _c.deepcopy(p0._p); p0._p.getparent().append(np_)
tf.paragraphs[1].runs[0].text = lines[1]; tf.paragraphs[2].runs[0].text = lines[2]
# elbow: stem (y 276) -> right to x 1014 -> down to box mid -> right into box
by_ = 276; my = ay + ah / 2; vx = 1014
c17 = S['Connector 17']; c17.begin_x, c17.begin_y, c17.end_x, c17.end_y = E(sx), E(by_), E(vx), E(by_)
for (a, b, c, d) in [(vx, by_, vx, my), (vx, my, bx, my)]:
    el = copy.deepcopy(c17._element); c17._element.addnext(el)
    ln = s.shapes[-1] if False else [x for x in s.shapes if x._element is el][0]
    ln.begin_x, ln.begin_y, ln.end_x, ln.end_y = E(a), E(b), E(c), E(d)
P.save('v7/Slide_Org_Structure.pptx'); print('ok')
