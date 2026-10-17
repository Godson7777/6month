"""McKinsey-style twin 100% stacked column: units vs value by crane class (Excel with formulas + 1-slide PPT)."""
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Font, PatternFill, Alignment
D = [('Light (≤ 8 tm)', 177, 40118000000), ('Small (8–25 tm)', 200, 77241000000),
     ('Medium (25–45 tm)', 196, 118328000000), ('Heavy (> 45 tm)', 65, 91154000000)]
COL = ['D9D9D9', 'BFBFBF', '595959', '000000']
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Units vs Value'
ws['A1'] = 'Heavy cranes are 10% of units but 28% of value'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = 'Indonesia knuckle boom crane imports, Jan 2023 – 14 Aug 2026'; ws['A2'].font = Font(italic=True, color='6B6B6B')
hdr = ['Class', 'Units', 'Value (IDR)', 'Units %', 'Value %', 'Value per unit (IDR)']
for c, h in enumerate(hdr, 1):
    x = ws.cell(4, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='000000')
for i, (k, u, v) in enumerate(D, 5):
    ws.cell(i, 1, k); ws.cell(i, 2, u); ws.cell(i, 3, v)
    ws.cell(i, 4, f'=B{i}/$B$9'); ws.cell(i, 5, f'=C{i}/$C$9'); ws.cell(i, 6, f'=C{i}/B{i}')
ws['A9'] = 'Total'; ws['B9'] = '=SUM(B5:B8)'; ws['C9'] = '=SUM(C5:C8)'; ws['D9'] = '=SUM(D5:D8)'; ws['E9'] = '=SUM(E5:E8)'; ws['F9'] = '=C9/B9'
for c in range(1, 7): ws.cell(9, c).font = Font(bold=True)
for r in range(5, 10):
    ws.cell(r, 2).number_format = '#,##0'; ws.cell(r, 3).number_format = '#,##0'; ws.cell(r, 6).number_format = '#,##0'
    ws.cell(r, 4).number_format = '0%'; ws.cell(r, 5).number_format = '0%'
ws['A11'] = 'Medium + Heavy'; ws['A11'].font = Font(bold=True)
ws['B11'] = '=B7+B8'; ws['C11'] = '=C7+C8'; ws['D11'] = '=D7+D8'; ws['E11'] = '=E7+E8'
ws['A12'] = 'Heavy value per unit vs Light'; ws['F12'] = '=F8/F5'; ws['F12'].number_format = '0.0"x"'
for c in 'BC': ws[f'{c}11'].number_format = '#,##0'
for c in 'DE': ws[f'{c}11'].number_format = '0%'
for c, w in zip('ABCDEF', [20, 10, 20, 10, 10, 22]): ws.column_dimensions[c].width = w
# chart helper table: rows = Units / Value, cols = classes (percent formulas)
ws['H4'] = ''; ws['H5'] = 'Units'; ws['H6'] = 'Value'
for j in range(4):
    col = openpyxl.utils.get_column_letter(9 + j)
    ws[f'{col}4'] = f'=A{5+j}'; ws[f'{col}5'] = f'=D{5+j}'; ws[f'{col}6'] = f'=E{5+j}'
    ws[f'{col}5'].number_format = ws[f'{col}6'].number_format = '0%'
ch = BarChart(); ch.type = 'col'; ch.grouping = 'percentStacked'; ch.overlap = 100; ch.gapWidth = 60
ch.title = 'Heavy cranes are 10% of units but 28% of value'
ch.add_data(Reference(ws, min_col=9, max_col=12, min_row=4, max_row=6), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=8, min_row=5, max_row=6))
for s, c in zip(ch.series, COL):
    s.graphicalProperties.solidFill = c; s.graphicalProperties.line.solidFill = 'FFFFFF'
    s.dLbls = DataLabelList(); s.dLbls.showVal = True; s.dLbls.numFmt = '0%'
ch.y_axis.majorGridlines = None; ch.y_axis.delete = True; ch.legend.position = 'r'; ch.width = 18; ch.height = 11
ws.add_chart(ch, 'H9')
wb.save('McK_Units_vs_Value.xlsx')

# ---- one slide, template font ----
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
FONT = 'Neue Haas Grotesk Text Pro'; E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string
P = Presentation(); P.slide_width, P.slide_height = 12192000, 6858000
s = P.slides.add_slide(P.slide_layouts[6])
def tx(x, y, w, h, t, sz, bold=False, color='000000', al=PP_ALIGN.LEFT):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = al; r = p.add_run(); r.text = t
    r.font.size = Pt(sz); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color); return b
tx(36, 40, 1200, 60, 'Heavy cranes are 10% of units but 28% of value', 32, True)
tx(36, 100, 1200, 26, 'Indonesia knuckle boom crane imports by class, Jan 2023 – 14 Aug 2026 (100% = 638 units; Rp 326,8 B)', 13, color='6B6B6B')
tu = sum(d[1] for d in D); tv = sum(d[2] for d in D)
cd = CategoryChartData(); cd.categories = ['Units', 'Value']
for k, u, v in D: cd.add_series(k, (u / tu, v / tv))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, E(36), E(140), E(760), E(540), cd); c = gf.chart
c.has_legend = True; c.legend.position = XL_LEGEND_POSITION.RIGHT; c.legend.include_in_layout = False; c.legend.font.size = Pt(12); c.legend.font.name = FONT
c.plots[0].gap_width = 60; c.plots[0].overlap = 100
for ser, col in zip(c.series, COL):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col); ser.format.line.color.rgb = rgb('FFFFFF')
    ser.data_labels.show_value = True; ser.data_labels.number_format = '0%'; ser.data_labels.number_format_is_linked = False
    ser.data_labels.font.size = Pt(13); ser.data_labels.font.bold = True; ser.data_labels.font.name = FONT
    ser.data_labels.font.color.rgb = rgb('000000' if col in ('D9D9D9', 'BFBFBF') else 'FFFFFF')
# series lines (McKinsey connectors)
from lxml import etree
from pptx.oxml.ns import qn
bc = c._chartSpace.find('.//' + qn('c:barChart'))
sl = etree.SubElement(bc, qn('c:serLines')); sp = etree.SubElement(sl, qn('c:spPr')); ln = etree.SubElement(sp, qn('a:ln')); ln.set('w', '9525')
sf = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr')).set('val', '7F7F7F')
axid = bc.findall(qn('c:axId'))
for a in axid: bc.remove(a); bc.append(a)
c.value_axis.visible = False; c.value_axis.has_major_gridlines = False
c.category_axis.tick_labels.font.size = Pt(14); c.category_axis.tick_labels.font.bold = True; c.category_axis.tick_labels.font.name = FONT
# takeaway column
tx(850, 170, 380, 30, 'Key takeaways', 16, True)
pts = [('10% → 28%', 'Heavy: 65 of 638 units, but Rp 91,2 B of Rp 326,8 B value'),
       ('41% → 64%', 'Medium + Heavy: 261 units, Rp 209,5 B'),
       ('6,2x', 'Heavy value per unit vs Light (Rp 1,40 B vs Rp 0,23 B)')]
for i, (big, small) in enumerate(pts):
    y = 215 + i * 150
    tx(850, y, 380, 56, big, 34, True); tx(850, y + 60, 380, 60, small, 13, color='404040')
tx(36, 690, 1200, 20, 'Source: Indonesia import records HS 84269100 + 84264900; value = import value (IDR).', 9, color='6B6B6B')
P.save('McK_Units_vs_Value.pptx'); print('ok', 91154000000/65/(40118000000/177))
