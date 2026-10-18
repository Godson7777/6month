"""STP Medium 25-45 tm: top-1 competitor model by VALUE per tier (Noir database), native table."""
import openpyxl, collections, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string
DB = 'data/Crane_Market_GVW_Pairing_v11_Noir.xlsx'
FX = 17803                                   # JISDOR 23 Sep 2026 (same as F.lli Ferrari deck)
CN = {'SANY PALFINGER', 'XCMG', 'BOB-LIFT', 'ZOOMLION'}   # 0% duty (ACFTA); others Europe 5%
R = [r for r in openpyxl.load_workbook(DB, data_only=True, read_only=True)['8. Pivot Source'].iter_rows(min_row=2, values_only=True) if r[2] == 'Medium']
def fam(b, m):          # merge variants: SPK32080B -> SPK32080, PK 41002 EH C -> PK 41002
    out = []
    for t in m[len(b):].strip().split():
        if re.search(r'\d', t): out.append(re.sub(r'(\d)[A-Z]+$', r'\1', t)); break
        out.append(t)
    return ' '.join(out)
def price(b, usd):      # highest import unit price x FX x (1+duty) x (1 + PPN 11% + PPh22 2.5%), round Rp 10 Jt
    return round(usd * FX * (1 + (0 if b in CN else .05)) * 1.135 / 1e7) * 1e7
tier = lambda tm: next(lo for lo in (25, 30, 35, 40) if lo < tm <= lo + 5)
CV = sum(r[10] for r in R); CU = sum(r[8] for r in R)
FER = {25: ('268 A4', 500e6), 30: ('FBR350R A4', 1090e6), 35: ('7441C*', 1240e6), 40: ('746 A4', 1420e6)}
NICE = lambda b: {'SANY PALFINGER': 'Sany Palfinger', 'PALFINGER': 'Palfinger', 'HIAB': 'HIAB', 'AMCO VEBA': 'Amco Veba', 'XCMG': 'XCMG'}.get(b, b.title())
ROWS = []
for t in (25, 30, 35, 40):
    rr = [r for r in R if tier(r[12]) == t]; tv = sum(r[10] for r in rr); tq = sum(r[8] for r in rr)
    V = collections.Counter(); Q = collections.Counter(); MX = collections.Counter()
    for r in rr:
        k = (r[0], fam(r[0], r[1])); V[k] += r[10]; Q[k] += r[8]; MX[k] = max(MX[k], r[9])
    k, v = V.most_common(1)[0]; p = price(k[0], MX[k]); fn, fp = FER[t]
    ROWS.append(dict(t=t, tq=tq, ms=tv / CV, brand=NICE(k[0]), type=k[1], q=Q[k], sh=v / tv, p=p, fn=fn, fp=fp, gap=fp / p - 1))
rp = lambda x: 'Rp ' + f'{x:,.0f}'.replace(',', '.')
pct = lambda x: f'{x:.0%}'
for r in ROWS: print(r)

P = Presentation(); P.slide_width, P.slide_height = 12192000, 6858000
s = P.slides.add_slide(P.slide_layouts[6])
FONT, SERIF = 'Arial', 'Georgia'
def tx(x, y, w, h, runs, sz, font=FONT, al=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    b = s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = al
    for t, bold, c, *it in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = font
        r.font.color.rgb = rgb(c); r.font.italic = bool(it and it[0])
        if spacing: r.font._element.set('spc', str(spacing))
    return b
def line(x1, y1, x2, y2, c='E6E8EB', w=1):
    l = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)); l.line.color.rgb = rgb(c); l.line.width = Pt(w)
# header
s.shapes.add_picture('assets/logos/triatra.png', E(56), E(18), height=E(24))
tx(700, 22, 524, 14, [('NEDP 6-MONTH REVIEW  ·  STP · MEDIUM CLASS 25–45 TM · 1 / 2', False, '697177')], 8, al=PP_ALIGN.RIGHT, spacing=100)
line(56, 50, 1224, 50)
tx(56, 70, 900, 20, [('STP · MEDIUM CLASS 25–45 TM · 1 / 2', True, '697177')], 14, spacing=200)
tx(56, 98, 1170, 40, [('Medium: target the value leaders with ', False, '11181C'), ('FERRARI', True, '076932', True), (' at 25–40 tm', False, '11181C')], 26, SERIF)
# STP boxes
mid = [r for r in ROWS if r['t'] in (30, 35)]
seg = (f"Medium class split into 4 tm tiers. 30–40 tm holds {sum(r['ms'] for r in mid):.0%} of Medium value "
       f"({sum(r['tq'] for r in mid)} of {CU} units).")
tgt = 'Top 1 model by value in each tier: ' + '; '.join(f"{r['brand']} {r['type']} ({r['t']}–{r['t'] + 5})" for r in ROWS) + '.'
win = [r for r in ROWS if r['gap'] < 0]; lose = [r for r in ROWS if r['gap'] > 0]
pos = (f"Ferrari is {min(-r['gap'] for r in win):.0%}–{max(-r['gap'] for r in win):.0%} cheaper than the value leader at 25–40 tm "
       f"({sum(r['ms'] for r in win):.0%} of Medium value). At 40–45 tm, 746 A4 is {lose[0]['gap']:.0%} above SPK42502.")
for i, (head, body, fill) in enumerate([('SEGMENTATION', seg, 'F9FAFB'), ('TARGETING', tgt, 'F9FAFB'), ('POSITIONING', pos, 'E1F9EC')]):
    x = 56 + i * 396
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(160), E(380), E(108)); b.adjustments[0] = 0.08
    b.fill.solid(); b.fill.fore_color.rgb = rgb(fill); b.line.fill.background(); b.shadow.inherit = False
    tx(x + 14, 172, 352, 14, [(head, True, '076932')], 9, spacing=150)
    tx(x + 14, 186, 352, 78, [(body, False, '11181C')], 12)
# table
cols = [74, 52, 74, 136, 116, 50, 60, 150, 100, 126, 230]
hdr1 = ['TM Class', 'Market Size', '', 'Top 1 Competitor (by value)', '', '', '', '', 'Challenger', '', '']
hdr2 = ['', 'Qty', '% of Medium', 'Brand', 'Type', 'Qty', 'Value share', 'Price/unit', 'Brand', 'Type', 'Price/Unit']
nr = 3 + len(ROWS)
tb = s.shapes.add_table(nr, len(cols), E(56), E(282), E(sum(cols)), E(24 + 26 + 26 + 64 * len(ROWS))).table
for j, w in enumerate(cols): tb.columns[j].width = E(w)
tb.rows[0].height = E(24); tb.rows[1].height = E(26); tb.rows[2].height = E(26)
for i in range(3, nr): tb.rows[i].height = E(64)
def cell(i, j, text, bold=False, fill='FFFFFF', color='11181C', sz=11):
    c = tb.cell(i, j); c.text = text; c.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.margin_left = c.margin_right = E(4); c.margin_top = c.margin_bottom = E(1)
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
    for p in c.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs: r.font.name = FONT; r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = rgb(color)
    return c
c = tb.cell(0, 0); c.merge(tb.cell(0, len(cols) - 1)); cell(0, 0, 'Medium Class Crane', True, '11181C', 'FFFFFF', 12)
for j, t in enumerate(hdr1): 
    if t or j in (0,): pass
cell(1, 0, 'TM Class', True, 'F0F2F4', sz=10); tb.cell(1, 0).merge(tb.cell(2, 0))
for (a, b_, t) in [(1, 2, 'Market Size'), (3, 7, 'Top 1 Competitor (by value)'), (8, 10, 'Challenger')]:
    tb.cell(1, a).merge(tb.cell(1, b_)); cell(1, a, t, True, 'F0F2F4', sz=10)
for j, t in enumerate(hdr2):
    if j: cell(2, j, t, True, 'F0F2F4', sz=10)
for i, r in enumerate(ROWS, start=3):
    gfill = 'E1F9EC' if r['gap'] < 0 else ('FFF7C2' if r['gap'] <= .30 else 'FEEBEC')
    vals = [f"{r['t']} – {r['t'] + 5}", str(r['tq']), pct(r['ms']), r['brand'], r['type'], str(r['q']), pct(r['sh']), rp(r['p'])]
    for j, v in enumerate(vals): cell(i, j, v, j in (0, 3))
    cell(i, 8, 'Ferrari', True, gfill); cell(i, 9, r['fn'], False, gfill); cell(i, 10, f"{rp(r['fp'])}  ({r['gap']:+.0%})".replace('-', '−'), False, gfill)
foot = ('Period Jan 2023 – 14 Aug 2026; source Crane Market GVW Pairing v11 (Noir). Tier by original max lifting moment; model variants merged (e.g. SPK32080 = SPK32080, B, C). '
        '% of Medium = tier import value / Medium import value. Top 1 = model with the highest import value in the tier; Share = model value / tier value. '
        'Competitor price = highest import unit price × Rp 17.803/USD (JISDOR 23 Sep 2026) × (1 + duty: 0% China ACFTA, 5% Europe) × (1 + PPN 11% + PPh 22 2.5%), before distributor margin. '
        'Ferrari = TSP price list Rev1 Jun 2026 less 13.6% TSP margin and 3% warranty; *7441C estimate. Fill: green = Ferrari cheaper, yellow = up to +30%, red = over +30%.')
tx(56, 642, 1168, 40, [(foot, True, '697177')], 7)
line(56, 688, 1224, 688)
tx(56, 694, 600, 12, [('NEDP MID YEAR REVIEW 2026  ·  TIBERIAS K. SIMU', True, '889096')], 7.5, spacing=150)
tx(1100, 694, 124, 12, [('12 / 18', True, '889096')], 7.5, al=PP_ALIGN.RIGHT)
P.save('STP_Medium_TopValue.pptx'); print('saved')
