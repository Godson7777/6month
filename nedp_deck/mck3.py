"""3-bar McKinsey 100% stacked chart (database: 764 units, value USD): Units | Value | Share, with series lines."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree
FONT='Neue Haas Grotesk Text Pro'; E=lambda v: Emu(round(v*9525)); rgb=RGBColor.from_string
CLS=['Light','Small','Medium','Heavy']; NM={'Light':'Light (≤ 8 tm)','Small':'Small (8–25 tm)','Medium':'Medium (25–45 tm)','Heavy':'Heavy (> 45 tm)'}
U={'Light':248,'Small':223,'Medium':218,'Heavy':75}  # database: units 2023-2026*
V={'Light':3019087,'Small':5037297,'Medium':7304431,'Heavy':5712582}  # database: value USD
COL={'Light':'C7EFD5','Small':'6FCF97','Medium':'1E9E57','Heavy':'0B4D2A'}
TU=sum(U.values()); TV=sum(V.values())
su={c:U[c]/TU for c in CLS}; sv={c:V[c]/TV for c in CLS}
rp=lambda v: f'${v/1e6:,.2f} M'.replace('.',',')
P=Presentation(); P.slide_width,P.slide_height=12192000,6858000
s=P.slides.add_slide(P.slide_layouts[6])
def tx(x,y,w,h,t,sz,b=False,c='000000'):
    tf=s.shapes.add_textbox(E(x),E(y),E(w),E(h)).text_frame; tf.word_wrap=True
    r=tf.paragraphs[0].add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.name=FONT; r.font.color.rgb=rgb(c)
tx(36,34,1208,56,f'Heavy is {su["Heavy"]:.0%} of units but {sv["Heavy"]:.0%} of value',30,True)
tx(36,92,1208,24,f'Crane imports by class, 2023 – 2026* (100% = {TU} units = {rp(TV)})',13,c='6B6B6B')
cats=['Units (qty)','Value (USD)','Share (units → value)']
cd=CategoryChartData(); cd.categories=cats
for c in CLS: cd.add_series(NM[c],[U[c],V[c]/1e6,V[c]/1e6])
gf=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100,E(36),E(130),E(1000),E(540),cd); ch=gf.chart
ch.has_legend=True; ch.legend.position=XL_LEGEND_POSITION.RIGHT; ch.legend.include_in_layout=False
ch.legend.font.size=Pt(12); ch.legend.font.name=FONT
pl=ch.plots[0]; pl.gap_width=60; pl.overlap=100
ch.value_axis.visible=False; ch.value_axis.has_major_gridlines=False
ta=ch.category_axis.tick_labels; ta.font.size=Pt(14); ta.font.bold=True; ta.font.name=FONT
lab=[lambda c:f'{U[c]} units\n{su[c]:.0%}', lambda c:f'{rp(V[c])}\n{sv[c]:.0%}', lambda c:f'{su[c]:.0%} → {sv[c]:.0%}']
for ser,c in zip(ch.series,CLS):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb=rgb(COL[c]); ser.format.line.color.rgb=rgb('FFFFFF')
    dl=ser.data_labels; dl.font.size=Pt(13); dl.font.bold=True; dl.font.name=FONT
    for k in range(3):
        tf=ser.points[k].data_label.text_frame; tf.text=lab[k](c)
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.bold=True; r.font.color.rgb=rgb('000000' if c in ('Light','Small') else 'FFFFFF')
# McKinsey connector lines = native series lines
bc=ch._chartSpace.find('.//'+qn('c:barChart'))
sl=etree.SubElement(bc,qn('c:serLines'))
sl.append(etree.fromstring('<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ln w="9525"><a:solidFill><a:srgbClr val="7F7F7F"/></a:solidFill></a:ln></c:spPr>'))
ax=bc.findall(qn('c:axId'))[0]; ax.addprevious(sl)
tx(36,676,1208,20,'*2026 = partial year (to Aug). Source: Indonesia import records, knuckle boom cranes; value in USD (M = million). Share = % of total units / % of total value.',9,c='6B6B6B')
P.save('McK_3bar_Units_Value_Share.pptx'); print(TU,TV,{c:(round(su[c]*100),round(sv[c]*100)) for c in CLS})
