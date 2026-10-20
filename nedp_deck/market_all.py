"""All-class crane market map (Noir, Hyva + Zoomlion excluded): tiers, top-1 model by value, avg landed price, cheapest challenger that beats top 1."""
import openpyxl, collections, re, json
DB = 'data/Crane_Market_GVW_Pairing_v11_Noir.xlsx'; FX = 17803
CN = {'SANY PALFINGER', 'XCMG', 'BOB-LIFT', 'ZOOMLION'}
EXCL = {'HYVA', 'ZOOMLION'}
R = [r for r in openpyxl.load_workbook(DB, data_only=True, read_only=True)['8. Pivot Source'].iter_rows(min_row=2, values_only=True)
     if r[0] and r[0] not in EXCL and r[8]]
def fam(b, m):
    out = []
    for t in m[len(b):].strip().split():
        if re.search(r'\d', t): out.append(re.sub(r'(\d)[A-Z]+$', r'\1', t)); break
        out.append(t)
    return ' '.join(out)
landed = lambda b, usd: usd * FX * (1 + (0 if b in CN else .05)) * 1.135
HI = {'Light': 8, 'Small': 25, 'Medium': 45, 'Heavy': 999}
def cls_of(tm):                      # AE rule: class by rounded tm (25.3 -> 25 -> Small)
    t = int(tm + .5); return 'Light' if t <= 8 else 'Small' if t <= 25 else 'Medium' if t <= 45 else 'Heavy'
def tier(cls, tm):
    tm = min(tm, HI[cls])
    if cls == 'Light': return (0, 5) if tm <= 5 else (5, 8)
    if cls == 'Small': return (8, 10) if tm <= 10 else next((lo, lo + 5) for lo in (10, 15, 20) if tm <= lo + 5)
    lo = int((tm - 1e-9) // 5 * 5); return (lo, lo + 5)
# F.lli Ferrari: cost = (K - J) x (1 - 13.6%), sheet 25.06; tm per flliferrari.it (class / model page) or name approach
wb = openpyxl.load_workbook('data/Price_List_Crane_Ferrari_Jun26_Rev1.xlsx', data_only=True)['25.06']
TM = {'315 A2': 1.5, '350 A2': 3.45, '543 A2': 4.3, '563 A2': 6, '563R A2': 6, '563 A3': 6, '563 A4': 6, '563R A4': 6,
      '6108 A2': 9.6, '6108 A3': 9.6, '6148 A3': 13.9, '5148 A4': 13.9, '178 A3': 15.8, '178 A4': 15.8, '6238 A3': 19.3, '6238 A4': 19.3,
      '6248 A3': 23, '6248 A4': 23, '258 A3': 25, '258R A3': 25, '258 A4': 25, '268 A3': 25.4, '268 A4': 25.4,
      '934 A4': 33, '934R A4': 33, 'FBR350 A4': 32.8, 'FBR350R A4': 32.8, 'FBR450 A4': 45.5, 'FBR450R A4': 45.5, 'FBR450R A6': 45.5,
      '746 A4': 44, '749R A4': 49, 'FBR600R A4': 58, 'FBR660R A4': 58.8, 'FBR660R A6': 58.8}
FER = []
for r in wb.iter_rows(min_row=3, values_only=True):
    if r[3] in TM and r[10]:
        FER.append(dict(brand='F.lli Ferrari', type=r[3], tm=TM[r[3]], p=(r[10] - r[9]) * (1 - .136), lead=r[6], src=r[4]))
# keep cheapest row per Ferrari model name (350 A2 Italy vs Malaysia)
fb = {}
for f in FER:
    if f['type'] not in fb or f['p'] < fb[f['type']]['p']: fb[f['type']] = f
FER = list(fb.values())
# competitor models: avg landed price = sum(value)/sum(qty)
M = collections.defaultdict(lambda: dict(q=0, v=0, tms=[]))
for r in R:
    k = (r[2], r[0], fam(r[0], r[1])); m = M[k]; m['q'] += r[8]; m['v'] += r[10]; m['tms'].append(r[12])
rows = []
for (cls, b, t), m in M.items():
    tm = max(m['tms']); rows.append(dict(cls=cls, brand=b, type=t, tm=tm, tier=tier(cls, tm), q=m['q'], usd=m['v'],
                                         p=landed(b, m['v'] / m['q']), amt=landed(b, m['v'])))
out = []
for cls in ('Light', 'Small', 'Medium', 'Heavy'):
    for tr in sorted({x['tier'] for x in rows if x['cls'] == cls}):
        T = [x for x in rows if x['cls'] == cls and x['tier'] == tr]; T.sort(key=lambda x: -x['usd'])
        top = T[0]; tv = sum(x['usd'] for x in T); tq = sum(x['q'] for x in T)
        cands = [dict(brand=x['brand'], type=x['type'], tm=x['tm'], p=x['p']) for x in T[1:]] + \
                [f for f in FER if cls_of(f['tm']) == cls and tier(cls, f['tm']) == tr]
        win = sorted([c for c in cands if c['p'] < top['p']], key=lambda c: c['p'])
        out.append(dict(cls=cls, tier=tr, units=tq, usd=tv, top=top, models=T, ch=win[0] if win else None,
                        fer=[f for f in FER if cls_of(f['tm']) == cls and tier(cls, f['tm']) == tr]))
json.dump(out, open('/tmp/claude-0/ps/market_all.json', 'w'), default=str, indent=1)
tot_q = sum(x['q'] for x in rows); tot_v = sum(x['amt'] for x in rows)
print(f"TOTAL {tot_q} units, landed Rp {tot_v/1e9:,.1f} M, import value Rp {sum(x['usd'] for x in rows)*FX/1e9:,.1f} M")
for o in out:
    t = o['top']; c = o['ch']
    print(f"{o['cls']:6s} {o['tier'][0]:>2}-{o['tier'][1]:<2} {o['units']:3d}u | TOP {t['brand'][:14]:14s} {t['type'][:14]:14s} {t['tm']:5.1f}tm {t['q']:3d}u Rp{t['p']/1e6:7,.0f}Jt | "
          + (f"CH {c['brand'][:14]:14s} {c['type'][:12]:12s} {c['tm']:5.1f}tm Rp{c['p']/1e6:7,.0f}Jt  -Rp{(t['p']-c['p'])/1e6:,.0f}Jt" if c else "CH none")
          + " | FER " + ', '.join(f"{f['type']} {f['p']/1e6:,.0f}" for f in o['fer']))
