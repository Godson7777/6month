"""Shared frame + drawing helpers for the crane deck (user's deck style, frame copied from 6M_V9 slide 14)."""
import copy
import crane_data as d
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
E = lambda v: Emu(round(v * 9525)); rgb = RGBColor.from_string; FONT = 'Neue Haas Grotesk Text Pro'
num = lambda x, n=1: f'{x:,.{n}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
COL = {'Light': 'CDEFD9', 'Small': '74CC95', 'Medium': '1E9E57', 'Heavy': '0B4D2A'}
DARK = {'Medium', 'Heavy'}; G = '1E9E57'; DG = '0B4D2A'; LG = 'E1F9EC'
LEG = {'Light': 'Light (≤ 8 tm)', 'Small': 'Small (8–25 tm)', 'Medium': 'Medium (25–45 tm)', 'Heavy': 'Heavy (> 45 tm)'}
class Deck:
    def __init__(self):
        self.P = Presentation('v7/6M_V9.pptx'); L = self.P.slides._sldIdLst
        for i, sid in reversed(list(enumerate(list(L)))):
            if i != 13: self.P.part.drop_rel(sid.rId); L.remove(sid)
        self.src = self.P.slides[0]; self.first = True
    def slide(self, title, q1, q2, head, parts, title_w=860):
        s = self.P.slides.add_slide(self.src.slide_layout)
        for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
        for sh in self.src.shapes:
            if sh.name in ('Date Placeholder 3', 'Slide Number Placeholder 5', 'Title 10', 'Group 31', 'Straight Connector 34', 'Table 2'):
                s.shapes._spTree.append(copy.deepcopy(sh._element))
        SH = {sh.name: sh for sh in s.shapes}
        grp = SH['Group 31']; xf = grp._element.find('.//' + qn('a:xfrm'))
        for tag in ('a:ext', 'a:chExt'): xf.find(qn(tag)).set('cx', str(E(1242 - 36.45)))
        for c in grp.shapes:
            ps = c.text_frame.paragraphs
            if 'Medium' in c.text_frame.text: set_runs(ps[0], title); c.width = E(title_w)
            else:
                set_runs(ps[0], q1); r0_ = ps[0].runs[0]._r
                br = etree.SubElement(ps[0]._p, qn('a:br')); ps[0]._p.remove(br); r0_.addnext(br)
                if r0_.find(qn('a:rPr')) is not None: br.append(copy.deepcopy(r0_.find(qn('a:rPr'))))
                r1_ = copy.deepcopy(r0_); br.addnext(r1_); r1_.find(qn('a:t')).text = q2
                c.left = E(900); c.width = E(342)
        t2 = SH['Table 2'].table
        set_runs(t2.cell(0, 0).text_frame.paragraphs[0], head)
        p1 = t2.cell(1, 0).text_frame.paragraphs[0]; r0 = p1.runs[0]
        for r in p1.runs[1:]: r._r.getparent().remove(r._r)
        r0.text = parts[0][0]; prev = r0._r
        for _ in parts[1:]:
            nr = copy.deepcopy(r0._r); prev.addnext(nr); prev = nr
        for r, (t, b) in zip(p1.runs, parts): r.text = t; r.font.bold = b
        return Canvas(s)
    def save(self, path):
        L = self.P.slides._sldIdLst; self.P.part.drop_rel(L[0].rId); L.remove(L[0])   # drop the frame source slide
        self.P.save(path)
def set_runs(par, text):
    rs = par.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
class Canvas:
    def __init__(self, s): self.s = s
    def tx(self, x, y, w, h, paras, sz=11, al=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        b = self.s.shapes.add_textbox(E(x), E(y), E(w), E(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for k, para in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = al
            for t, bold, c, *z in para:
                r = p.add_run(); r.text = t; r.font.size = Pt(z[0] if z else sz); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(c)
        return b
    def box(self, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1):
        b = self.s.shapes.add_shape(shape, E(x), E(y), E(w), E(h)); b.shadow.inherit = False
        if fill: b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
        else: b.fill.background()
        if line: b.line.color.rgb = rgb(line); b.line.width = Pt(lw)
        else: b.line.fill.background()
        return b
    def line(self, x1, y1, x2, y2, color='7F7F7F', w=1):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2)); c.line.color.rgb = rgb(color); c.line.width = Pt(w); return c
    def chart(self, ctype, x, y, w, h, cats, series, gap=60, serlines=True):
        cd = CategoryChartData(); cd.categories = cats
        for name, vals in series: cd.add_series(name, vals, number_format='0')
        ch = self.s.shapes.add_chart(ctype, E(x), E(y), E(w), E(h), cd).chart
        ch.has_legend = False; ch.has_title = False; ch.plots[0].gap_width = gap; ch.plots[0].overlap = 100
        if serlines:
            bc = ch._chartSpace.find('.//' + qn('c:barChart'))
            sl = etree.SubElement(bc, qn('c:serLines')); bc.remove(sl); bc.find(qn('c:overlap')).addnext(sl)
            ln = etree.SubElement(etree.SubElement(sl, qn('c:spPr')), qn('a:ln'), w='6350'); etree.SubElement(etree.SubElement(ln, qn('a:solidFill')), qn('a:srgbClr'), val='808080')
        va = ch.value_axis; va.visible = False; va.has_major_gridlines = False
        ca = ch.category_axis; ca.tick_labels.font.size = Pt(10.5); ca.tick_labels.font.name = FONT; ca.format.line.color.rgb = rgb('404040')
        return ch
    @staticmethod
    def manual(ch, PX, PY, PW, PH):
        pa = ch._chartSpace.find('.//' + qn('c:plotArea')); lay = pa.find(qn('c:layout'))
        if lay is None: lay = etree.SubElement(pa, qn('c:layout')); pa.remove(lay); pa.insert(0, lay)
        lay.append(etree.fromstring(f'<c:manualLayout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                                    f'<c:x val="{PX}"/><c:y val="{PY}"/><c:w val="{PW}"/><c:h val="{PH}"/></c:manualLayout>'))
    @staticmethod
    def point_label(se, j, lines, color, sz=10.5):
        dl = se.points[j].data_label; dl.position = XL_LABEL_POSITION.CENTER; tf = dl.text_frame
        for li, part in enumerate(lines):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            rr = p.add_run(); rr.text = part; rr.font.size = Pt(sz); rr.font.bold = True; rr.font.name = FONT; rr.font.color.rgb = rgb(color)
    @staticmethod
    def point_fmt(se, j, fmt, color, sz=10, bold=True):
        """Native data label (value + number format with literal text): font stays editable in PowerPoint."""
        dl = se.points[j].data_label; dl.position = XL_LABEL_POSITION.CENTER; dl.show_value = True
        dl.font.size = Pt(sz); dl.font.bold = bold; dl.font.name = FONT; dl.font.color.rgb = rgb(color)
        el = dl._dLbl
        for old in el.findall(qn('c:numFmt')): el.remove(old)
        nf = etree.SubElement(el, qn('c:numFmt')); el.remove(nf); nf.set('formatCode', fmt); nf.set('sourceLinked', '0')
        anchor = el.find(qn('c:spPr')) if el.find(qn('c:spPr')) is not None else el.find(qn('c:txPr'))
        anchor.addprevious(nf)
    @staticmethod
    def hide_label(se, j):
        se.data_labels  # make sure c:dLbls exists
        se._element.find(qn('c:dLbls')).insert(0, etree.fromstring(
            f'<c:dLbl xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:idx val="{j}"/><c:delete val="1"/></c:dLbl>'))
    def legend(self, x, y, step, items=None):
        items = items or [(COL[c], LEG[c]) for c in d.CLS]
        for i, (col, lab) in enumerate(items):
            self.box(x + i * step, y + 2, 10, 10, col, 'BFBFBF' if col in ('FFFFFF', 'F2F2F2') else None, lw=.5)
            self.tx(x + i * step + 14, y, step - 14, 14, [[(lab, False, '000000', 10.5)]])
    def divider(self, x, y1, y2):
        self.line(x, y1, x, y2, G, 2.25)
        o = self.box(x - 14, (y1 + y2) / 2 - 14, 28, 28, G, None, MSO_SHAPE.OVAL); tf = o.text_frame
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = '>'; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = rgb('FFFFFF'); r.font.name = FONT
    def callout(self, x, yc, w, big, rest, color):
        h = 28 + 13.5 * len(rest); self.box(x, yc - h / 2, w, h, 'FFFFFF', 'BFBFBF', lw=.75)
        self.tx(x + 3, yc - h / 2 + 4, w - 6, h - 6, [[(big, True, color, 12)]] + [[(t, False, '000000', 8.5)] for t in rest], al=PP_ALIGN.CENTER)
    def mix_panel(self, MX, MY, MW, MHh, cu, cv, left_x, right_x, cw=82):
        """Units / Value / Units->value 100% columns with braces + callouts. cu, cv: dict class -> units / value (IDR)."""
        TU = sum(cu.values()); TV = sum(cv.values())
        QX, QY, QW, QH = .02, .02, .96, .89
        ch = self.chart(XL_CHART_TYPE.COLUMN_STACKED_100, MX, MY, MW, MHh, ['Units (qty)', 'Value (IDR)', 'Units → value'],
                        [(LEG[c], [cu[c] / TU, cv[c] / TV, cv[c] / TV]) for c in d.CLS], gap=25)
        self.manual(ch, QX, QY, QW, QH)
        for k, c in enumerate(d.CLS):
            se = ch.plots[0].series[k]; se.format.fill.solid(); se.format.fill.fore_color.rgb = rgb(COL[c])
            fc = 'FFFFFF' if c in DARK else '000000'
            for j in range(3):
                sh = cu[c] / TU if j == 0 else cv[c] / TV
                if sh < .045: self.hide_label(se, j); continue
                lit = [f'{cu[c]} units · ', f'Rp {num(cv[c] / 1e9)} Bio · ', f'{cu[c] / TU:.0%} → '][j]
                self.point_fmt(se, j, f'"{lit}"0%', fc, 9.5 if j == 1 else 10)
        qx0, qy0, qw, qh = MX + MW * QX, MY + MHh * QY, MW * QW, MHh * QH; qcat = qw / 3; bw = qcat / 1.25
        ub_l = qx0 + qcat * .5 - bw / 2; sb_r = qx0 + qcat * 2.5 + bw / 2
        ls_u = (cu['Light'] + cu['Small']) / TU; mh_v = (cv['Medium'] + cv['Heavy']) / TV
        y_u = qy0 + qh * (1 - ls_u); y_v = qy0 + qh * mh_v; ybot = qy0 + qh
        self.box(ub_l - 18, qy0, 12, y_u - qy0 - 3, None, G, MSO_SHAPE.LEFT_BRACE, 1.5)
        self.box(ub_l - 18, y_u + 3, 12, ybot - y_u - 3, None, '000000', MSO_SHAPE.LEFT_BRACE, 1.5)
        self.box(sb_r + 6, qy0, 12, y_v - qy0 - 3, None, G, MSO_SHAPE.RIGHT_BRACE, 1.5)
        self.box(sb_r + 6, y_v + 3, 12, ybot - y_v - 3, None, '000000', MSO_SHAPE.RIGHT_BRACE, 1.5)
        vmh = cv['Medium'] + cv['Heavy']; vls = cv['Light'] + cv['Small']
        self.callout(left_x, (qy0 + y_u) / 2, cw, f'{1 - ls_u:.0%}', ['of total units', 'from Medium', '& Heavy'], G)
        self.callout(left_x, (y_u + ybot) / 2, cw, f'{ls_u:.0%}', ['of total units', 'from Light', '& Small'], '000000')
        self.callout(right_x, (qy0 + y_v) / 2, cw, f'{mh_v:.0%}', ['of total value', f'(Rp {num(vmh / 1e9)} Bio)', 'from Medium', '& Heavy'], G)
        self.callout(right_x, (y_v + ybot) / 2, cw, f'{1 - mh_v:.0%}', ['of total value', f'(Rp {num(vls / 1e9)} Bio)', 'from Light', '& Small'], '000000')
        return ls_u, mh_v
SOURCE_NOTE = ('Source: Indonesia import records (Crane Market GVW Pairing v11, Noir), 2023 – 2026; 2026 is a partial year. Folding cranes only; excludes the Noir "Excluded Rows", Hyva, Zoomlion, '
               'HIAB XG grabbers, Tadano TM-Z telescopic and 2 invalid-price records. Value = import value before tax × Rp 17.803/USD (JISDOR 23 Sep 2026).')
